from datetime import date, datetime, timezone
import csv
import io

import pytest

from api.services.territory_manual_pack import build_pack, qualifies, write_pack

NOW = datetime(2026, 9, 13, tzinfo=timezone.utc)
SOURCE = {'status': 'completed', 'completed_at': NOW.isoformat(), 'source_checksum_sha256': 'a' * 64}


def row(i, **updates):
    return {'id': str(i), 'provider_id': f'p{i}', 'name': f'Site {i}',
            'status': 'ACTIVE', 'region': 'West Midlands', 'overall_rating': 'Requires improvement',
            'signal_checked_at': NOW, **updates}


def pack(rows, **overrides):
    args = dict(region='West Midlands', buyer='requires_improvement', service_type='',
                observed_at=NOW, provider_names={f'p{i}': f'Provider {i}' for i in range(40)},
                source_evidence=SOURCE, order_reference='rehearsal')
    args.update(overrides)
    return build_pack(rows, **args)


def test_duplicate_locations_do_not_inflate_organisation_shortlist():
    result = pack([row(i) for i in range(30)] + [row(100, provider_id='p0')])
    assert len(result['locations']) == 31
    assert len(result['shortlist']) == result['provider_count'] == 30
    assert result['shortlist'][0]['provider_id'] == 'p0'
    assert result['shortlist'][0]['matching_locations'] == 2
    assert len({r['provider_id'] for r in result['shortlist']}) == 30


def test_unknown_missing_duplicate_and_wrong_scope_fail_safely():
    with pytest.raises(ValueError):
        pack([row(1), row(1)])
    with pytest.raises(ValueError):
        pack([row(1, provider_id='')])
    with pytest.raises(ValueError):
        pack([row(1)], buyer='everyone')
    assert pack([row(1, region='London'), row(2, status='INACTIVE')])['locations'] == []


def test_source_failure_is_not_hidden_by_recent_inspection_or_extract():
    result = pack([row(i) for i in range(30)], source_evidence={'status': 'running'})
    assert any('reconciliation' in b for b in result['blockers'])
    assert result['status'] == 'DRAFT - NOT FOR CUSTOMER DELIVERY'
    result = pack([row(i) for i in range(30)], source_evidence={**SOURCE, 'completed_at': '2026-08-01T00:00:00+00:00'})
    assert any('freshness' in b for b in result['blockers'])


def test_no_provider_name_substitution_and_small_population_blocks():
    result = pack([row(1)], provider_names={})
    assert result['shortlist'][0]['provider_name'] == ''
    assert len(result['blockers']) == 2


def test_buyer_population_boundaries_and_missing_inspection():
    assert qualifies(row(1, registration_date='2026-09-13'), 'new_90', NOW.date())
    assert not qualifies(row(1, registration_date='2026-09-14'), 'new_90', NOW.date())
    assert qualifies(row(1, last_inspection_date=None), 'stale_inspection', NOW.date())
    assert not qualifies(row(1, last_inspection_date=date(2026, 1, 1)), 'stale_inspection', NOW.date())
    result = pack([row(1)], buyer='stale_inspection')
    assert result['shortlist'][0]['published_fact'] == 'No inspection date is recorded in this source'


def test_export_contains_complete_dataset_workbook_and_separate_brief(tmp_path):
    from openpyxl import load_workbook
    result = pack([row(i) for i in range(30)] + [row(100, provider_id='p0', name='=HYPERLINK("bad")')])
    out = tmp_path / 'pack'
    evidence = write_pack(result, out)
    assert set(evidence['files']) == {'brief.pdf', 'territory.csv', 'shortlist.csv', 'territory.xlsx'}
    wb = load_workbook(out / 'territory.xlsx')
    assert wb['Territory locations'].max_row == 32
    assert wb['Shortlist'].max_row == 31
    assert not any(c.data_type == 'f' for ws in wb for r in ws for c in r)
    rows = list(csv.DictReader(io.StringIO((out / 'territory.csv').read_text(encoding='utf-8-sig'))))
    assert len(rows) == 31
    assert next(r for r in rows if r['location_id'] == '100')['location_name'].startswith("'=")
    with pytest.raises(FileExistsError):
        write_pack(result, out)
