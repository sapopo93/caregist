from pathlib import Path
import json, hashlib, shutil, textwrap
from datetime import date, timedelta
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, Rectangle
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent
FIG = OUT / 'figures'
FIG.mkdir(parents=True, exist_ok=True)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':11,'axes.spines.top':False,'axes.spines.right':False})
NAVY='#203B50'; TEAL='#24756F'; GOLD='#A87821'; RED='#A54444'; GREY='#58636D'

def save(fig, name):
    fig.savefig(FIG/name, dpi=220, bbox_inches='tight', facecolor='white')
    plt.close(fig)

fig, ax=plt.subplots(figsize=(10,4.0)); ax.set(xlim=(0,10),ylim=(0,4.2)); ax.axis('off')
labels=[('1 Enquiry received','Actual inbox receipt'),('2 Qualify','Buyer decision and deadline'),('3 Check feasibility','Source, criteria and capacity'),('4 Agree scope','One accepted version'),('5 Verify payment','Merchant match to order'),('6 Produce','PDF, CSV and manifest'),('7 Independent review','PASS on exact final files'),('8 Release','Approved recipient and route'),('9 Confirm access','Buyer opens all three files')]
pos=[(0,3),(3.5,3),(7,3),(7,1.65),(3.5,1.65),(0,1.65),(0,.3),(3.5,.3),(7,.3)]
for i,((title,sub),(x,y)) in enumerate(zip(labels,pos)):
    ax.add_patch(Rectangle((x,y),3,.9,facecolor='#EFF4F6',edgecolor='#B5C1CA',lw=.8))
    ax.text(x+.14,y+.58,title,color=NAVY,fontsize=12,fontweight='bold')
    ax.text(x+.14,y+.23,sub,color=GREY,fontsize=9.5)
    if i<8:
        nx,ny=pos[i+1]
        if y==ny: start=(x+3,y+.45) if nx>x else (x,y+.45); end=(nx,ny+.45) if nx>x else (nx+3,ny+.45)
        else: start=(x+1.5,y); end=(nx+1.5,ny+.9)
        ax.add_patch(FancyArrowPatch(start,end,arrowstyle='-|>',mutation_scale=12,color=TEAL,lw=1.4))
save(fig,'journey.png')

fig,ax=plt.subplots(figsize=(10,2.8))
names=['Feedback schedule\nwith follow-ups','Proposed discovery day','Proposed production day']
parts=[[1.5,2.333,1.5],[1.5,2,1.5],[1.5,1,2.5]]
left=[0,0,0]
for j,(name,col) in enumerate([('Operations, handoffs and log',GREY),('Prospecting and follow-ups',GOLD),('Production or discovery',TEAL)]):
    vals=[r[j] for r in parts]; ax.barh(range(3),vals,left=left,color=col,label=name,height=.52)
    for i,v in enumerate(vals):ax.text(left[i]+v/2,i,f'{v:.2f}h' if i==0 and j==1 else f'{v:g}h',ha='center',va='center',color='white',fontsize=11)
    left=[a+b for a,b in zip(left,vals)]
ax.axvline(5,color=RED,linestyle='--',lw=1.5); ax.text(5.06,-.45,'5h daily limit',color=RED,fontsize=10)
ax.set(yticks=range(3),yticklabels=names,xlim=(0,6),xlabel='Operator hours per day'); ax.invert_yaxis()
handles,legend_labels=ax.get_legend_handles_labels()
ax.legend(handles,['Operations / handoffs / log','Prospecting + follow-ups','Production / discovery'],loc='upper center',bbox_to_anchor=(.5,-.24),ncol=3,frameon=False,fontsize=9.6)
ax.spines['left'].set_visible(False); ax.tick_params(axis='y',length=0)
save(fig,'capacity.png')

fig,ax=plt.subplots(figsize=(10,3.8))
for a,b in [(5,7),(12,14)]:ax.axvspan(a-.15,b-.15,color='#F0F0F0',zorder=0)
for y,vals,col in [(1,[0,1,1,4,4],TEAL),(0,[0,3,4,9,10],GOLD)]:
    ax.plot([0,max(vals)],[y,y],color=col,lw=2)
    titles=['Enquiry','Scope + source','Paid','Release','Access confirmed']
    for i,(v,t) in enumerate(zip(vals,titles)):
        dy=.12 if i%2==0 else -.13
        if y==1 and i==2:dy=-.28
        if y==1 and i==4:dy=-.28
        ax.scatter(v,y,color=col,s=45,zorder=3); ax.text(v,y+dy,t,ha='center',va='bottom' if dy>0 else 'top',fontsize=9.3)
ax.set(yticks=[0,1],yticklabels=['Late payment / later access','Prompt complete enquiry'],ylim=(-.6,1.65),xlim=(-1,14))
ax.set_xticks(range(0,12));ax.set_xticklabels([(date(2026,10,26)+timedelta(days=i)).strftime('%d %b') for i in range(12)],rotation=35,ha='right',fontsize=9)
ax.text(11.3,1,'4 calendar days',color=TEAL,fontsize=11)
ax.text(11.3,0,'10 calendar days',color=GOLD,fontsize=11)
ax.spines['left'].set_visible(False);ax.tick_params(axis='y',length=0)
save(fig,'clocks.png')

fig,axs=plt.subplots(1,2,figsize=(10,4.1),gridspec_kw={'width_ratios':[1,1.35]})
hrs=[5,8,8.5,9.1]; cont=[745-h*30-25 for h in hrs]
axs[0].bar([str(h) for h in hrs],cont,color=[GREY,TEAL,TEAL,GOLD]);axs[0].set(xlabel='All-in hours per Brief',ylabel='Illustrative contribution (£)',title='Per-order model')
for i,v in enumerate(cont):axs[0].text(i,v+8,f'£{v:.0f}',ha='center',fontsize=10)
axs[0].set_ylim(0,650)
orders=[0,1,2,3]; losses=[n*745-3000-510-90-n*45-n*25 for n in orders]
axs[1].bar([str(n) for n in orders],losses,color=RED);axs[1].set(xlabel='Paid orders (not forecasts)',ylabel='Illustrative pilot result (£)',title='Whole-pilot model')
for i,v in enumerate(losses):axs[1].text(i,v-100,f'−£{abs(v):,}',ha='center',fontsize=10)
axs[1].set_ylim(-4100,0); fig.tight_layout(w_pad=2)
save(fig,'economics.png')

fig,ax=plt.subplots(figsize=(10,2.5));ax.axis('off');ax.set(xlim=(0,10),ylim=(0,2.5))
steps=[('BEFORE PAYMENT','G0–G5 readiness\n+ accepted order scope'),('PAYMENT','Verified merchant status\nMatched order, £745 GBP'),('BEFORE RELEASE','G6 exact-file review\n+ authorised recipient'),('AFTER RELEASE','Access confirmation\nthen product acceptance')]
for i,(title,sub) in enumerate(steps):
    x=i*2.55
    ax.text(x,.0 if False else 2.15,title,color=NAVY,fontsize=11,fontweight='bold')
    ax.text(x,1.6,sub,color=GREY,fontsize=10.1,linespacing=1.5,va='top')
    if i<3:ax.add_patch(FancyArrowPatch((x+1.85,.65),(x+2.42,.65),arrowstyle='-|>',color=TEAL,mutation_scale=13,lw=1.6))
ax.text(0,.12,'G3 applies to authorised marketing; customer-service messages use their separately approved route.',fontsize=9.8,color=GREY)
save(fig,'gates.png')

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
sec.header_distance=Inches(.27);sec.footer_distance=Inches(.27)
for n in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[n];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    if st.element.pPr is not None:
        for el in list(st.element.pPr):
            if el.tag==qn('w:pBdr'):st.element.pPr.remove(el)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.06
for n,s in [('Title',27),('Heading 1',20),('Heading 2',13),('Heading 3',11.5)]:
    doc.styles[n].font.size=Pt(s);doc.styles[n].paragraph_format.space_before=Pt(9);doc.styles[n].paragraph_format.space_after=Pt(7)
doc.styles['Title'].paragraph_format.space_before=Pt(0)
doc.styles['Subtitle'].font.size=Pt(12)
hp=sec.header.paragraphs[0];hp.text='CAREGIST  |  CUSTOMER JOURNEY AND OPERATING PLAN';hp.runs[0].font.size=Pt(8);hp.runs[0].font.color.rgb=RGBColor(0,0,0)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('11 October 2026  •  Planning document  |  ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');fp._p.append(fld)
doc.core_properties.title='CareGist customer journey and 30 day operating plan'
doc.core_properties.subject='Consolidated source-checked plan for 12 October to 10 November 2026'
doc.core_properties.author='CareGist planning review'
source=[]; page_num=0

def page(title, intro=None):
    global page_num
    page_num+=1
    pp=doc.add_heading(title,0 if page_num==1 else 1)
    pp.paragraph_format.page_break_before=page_num>1
    if pp._p.pPr is not None:
        for el in list(pp._p.pPr):
            if el.tag==qn('w:pBdr'):pp._p.pPr.remove(el)
    source.append('\n# '+title+'\n')
    if intro:p(intro)
def p(text, bold=False):
    pp=doc.add_paragraph();r=pp.add_run(text);r.bold=bold;source.append(text+'\n');return pp
def h(text):doc.add_heading(text,2);source.append('\n## '+text+'\n')
def pic(name,caption):
    pp=doc.add_paragraph();pp.add_run().add_picture(str(FIG/name),width=Inches(6.95))
    pp.paragraph_format.space_after=Pt(3)
    pp=doc.add_paragraph(caption);pp.runs[0].font.size=Pt(9);pp.paragraph_format.space_after=Pt(8)
    source.append('[Figure: '+caption+']\n')
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    for row in rows:
        c=t.add_row().cells
        for i,text in enumerate(row):c[i].text=str(text)
    for ri,row in enumerate(t.rows):
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge,v in [('top','80'),('bottom','80'),('left','100'),('right','100')]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),v);el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'203B50' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(0);pp.paragraph_format.line_spacing=1.03
                if i==0 and widths[i]<1.3:pp.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in pp.runs:r.font.size=Pt(10.1);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
        trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(0);pp.paragraph_format.line_spacing=.4
    source.append('| '+' | '.join(headers)+' |\n')
    for row in rows:source.append('| '+' | '.join(str(x) for x in row)+' |\n')

page('CareGist customer journey and 30 day operating plan')
p('12 October–10 November 2026 • Prepared for Henry and Sibongile • Final consolidated planning edition',True)
p('Prove one £745 Territory Opportunity Brief through a staffed, controlled service. Aim for two paid, buyer-accepted Briefs; a third is conditional on proven capacity. Zero sales is possible. The first milestone is a complete rehearsal, not a payment request.')
pic('journey.png','Figure 1. The service begins with a received enquiry and reaches receipt only when the agreed recipient confirms all files open.')
h('What this plan gives you')
p('A complete enquiry-to-receipt process; corrected feedback; named decisions and evidence gates; a five-hour daily schedule; a dated 30-day work programme; approved-use draft messages; and transparent costs and measures.')
table(['Read this to…','Pages'],[('Understand the corrections and ownership','2–4'),('Run the customer journey and gates','5–9'),('Allocate time and execute the next 30 days','10–14'),('Manage outreach, cost and learning','15–17'),('Use message drafts and check sources','18–20')],[5.85,1.1])
p('Status: planning complete; commercial readiness remains unverified. This document does not authorise outreach, payment, refund, deployment or gate activation. Repository rules, approved terms and actual gate receipts take precedence. Use this as the consolidated proposed schedule; earlier plans remain evidence, not competing schedules.',True)

page('Feedback checked and corrected','The Claude Operating Bible was read in full. Its practical suggestions are retained where supported. The table below records the consequential corrections; a proposed fix is not a completed operational fix.')
table(['Feedback reference','Finding','Final treatment'],[
('D1 + D3 / B3–B4','Henry is both generator operator and proposed independent reviewer.','Treat Henry and Sibongile as producers if they generate or curate. Name a separate competent reviewer before paid commitments; reserve 1.5h provisionally, with more time if required.'),
('F17 / G5–G6','Requiring G6 before payment is circular; rehearsing only before release is too late.','G0–G5 readiness before a paid commitment/payment request; G6 checks paid order and final files before release. Mon 26 Oct is a planning earliest date only.'),
('D4 / B1 / F6','Local Brief checkout uses separate flags and does not call the Radar readiness helper.','This narrow code distinction is supported. It does not approve a manual current-status route or waive project source/export/commercial gates.'),
('D4 / F2 / F14','The existing generator selects in-window events; a weekly CSV alone is not proven sufficient for all promised fields.','Define and independently prove the selected source, factual predicates, ranking and output method. Otherwise use only a proven supported scope or refuse.'),
('D1–D11','“Silence means default stands” converts missing decisions into approval.','All blocking decisions require a named, explicit recorded decision. Open decisions remain open.'),
('Outreach / capacity','Five scheduled hours plus 20 minutes of follow-ups exceeds five hours.','Fit sending, research and follow-ups inside the allocated block. Reduce new first touches when service work requires capacity; no assumed overtime.'),
('G1 / T6','A fixed 7 October edition and flags-off check do not establish source readiness.','Recheck approved edition, retrieval time/hash, freshness, completeness and material exceptions for each scope and order.'),
('F7 / T6','The scope template omits the later source-path check and risks implying payment resets the clock.','Use the later of written scope confirmation and source check; agree an explicit calendar date. Late payment needs written date re-agreement.')
],[1.12,2.62,3.21])
p('The principal growth product manager review was performed by an invited AI subagent. It is analytical challenge of this plan, not an external human appointment or the prescribed commercial approval.')

page('Remaining feedback resolutions')
table(['Reference','Evidence check','Action in this final plan'],[
('B2 / cohort','Tracker: 20 rows; 16 FIT, 4 REVIEW; 17 NOT_STARTED, 3 CONTACTED. Fulcrum sample request and deferral recorded.','Prioritise relevant warm history after verifying the correspondence and suppression. No sale or new-contact permission inferred. The 11 VERIFY / 2 SKIP breakdown is not in this tracker schema.'),
('D5–D6 / B5','Stripe invoice and restricted cloud folder are plausible proposed routes, not proven operating capabilities.','Select actual merchant/storage configuration and provider. Test order matching, recovery, recipient denial, revocation and actual expiry capability before approval.'),
('B6 / D10','Training guide refers to South Africa; actual operator location is unconfirmed.','Agree UK-time coverage explicitly. If based in South Africa, the offset changes from +1h to +2h when UK clocks change on 25 Oct.'),
('F8 / journey','Prepared, released, usable access and product acceptance are distinct.','Separate statuses, timestamps and evidence. Tracking pixels, redirects, silence and “fulfilled” database labels do not prove receipt or acceptance.'),
('F9 / B7','Published terms include older products and omit the proposed Digest.','Sell only the Brief under approved operative terms and accepted scope. Do not declare that a new scope automatically overrides the terms. Hold Digest sales.'),
('F10','Prior Excel-workbook allegation was withdrawn after page inspection.','Pilot inventory is PDF, CSV and source/licence manifest. Do not promise a native Excel workbook.'),
('F11','Screening import expects screening fields, not a revenue tracker.','Use one restricted register until contact/deal/task migration, reconciliation, permissions and suppression are proven.'),
('F12–F13','Local and public releases differ. leads/README is already labelled research-only and superseded.','Pin exact deployed source and recheck affected findings. Keep historical leads research-only; avoid claiming a missing warning already exists.'),
('F15–F16','Auto-reply is unproven. Poll scheduling intervention does not guarantee readiness.','Test inbox acknowledgement. Poll evidence requires at least 24 total valid polls and ≥90% success in the rolling window, with other applicable checks intact.'),
('B8–B9 / T9–T13','Samples, expiry wording, remedy dates and “more useful” claims require evidence.','Create three reviewed samples alongside a full-sized rehearsal. Use tested access wording, realistic response dates and factual offer language. One consolidated schedule.')
],[1.08,2.67,3.2])
p('Feedback checking establishes a defensible operating design. It does not establish that staff accounts, source files, payment configuration, private delivery or commercial gates are ready today.')

page('People and decisions before commitments')
table(['Role','Responsibility','Capacity / decision'],[
('Sibongile','Own every enquiry, qualification, scope draft, curation, customer updates, access checks and recovery.','Confirmed input: 25h/week. Four core weeks =100h. No extra daily hour assumed.'),
('Henry','Commercial authority, terms/tax decisions, merchant operations, technical/source work, backup and pilot decision.','Confirm actual founder/engineering/backup calendar. Up to 17 engineering hours is an allocation proposal, not an estimate that all fixes fit.'),
('Independent reviewer','Check every shortlisted entry and exact rendered files; issue PASS or defects; recheck corrections.','Named person, competence and independence must be recorded. Budget 1.5h/Brief provisionally. Neither Henry nor Sibongile reviews a pack they produced.'),
('Backup','Cover inbox, deadlines and recovery with approved access and handover.','Henry is proposed; confirm availability. If unavailable, pause new commitments and give existing enquiries a next-update date.')
],[1.15,3.0,2.8])
h('Monday decisions with explicit records')
p('At 09:00 on 12 October, record reviewer and backup names; generator/source owner; supported scope and source method; payment route; actual storage provider and access method; cost ceiling and acceptable loss; first segment and region; shortlist target; UK coverage times; and acknowledgement wording/configuration. Record decision-maker, date, evidence, limits and follow-up owner. Thirty minutes may not resolve all eleven; unresolved items retain an owner and deadline.')
h('One offer')
p('A £745 one-off research pack for a supplier making an adult social care territory decision in England. Start with 25 distinct CQC provider IDs within the published 25–50 range, only after eligibility is proven. Confirm tax treatment before quoting. Use one agreed region or local authority and explicit service taxonomy, selection predicates, exclusions and any commercial-group roll-up.')
p('Deliver an executive PDF (target three to five pages, with a labelled appendix if needed), CRM-ready CSV and source/licence manifest. Each entry has source-linked facts, provider and relevant location IDs, a fit explanation, a qualification question and uncertainty. No promised contact enrichment, budget, staffing shortage, future inspection, compliance opinion, exclusive intelligence or guaranteed sale.')
p('Preferred first cohort: the verified relevant compliance/quality consultancies in the tracker, then training suppliers in the same region. This changes the earlier training-first hypothesis; compare actual buyer response rather than assuming one segment is better. West Midlands is a working example until selected.')
p('Terms used here: a provider is a registered organisation; a location is a registered care site/service. A gate is a check required before proceeding. A gate receipt is the evidence supporting its approval. A file hash identifies the exact version that was reviewed.')

page('Enquiry to verified payment','Every handoff must leave an owner, next action, next-update date and observable evidence. Proposed internal response targets become customer promises only after staffing and approval.')
table(['Stage','Action and owner','Evidence required to advance'],[
('1 Received','Sibongile monitors actual inbox receipt, assigns enquiry ID and acknowledges by next working day.','Original message, received timestamp, reference, owner and actual acknowledgement. Mailto click/composer is intent only.'),
('2 Qualified','Ask only missing questions: what the buyer sells, target providers, decision, deadline, exclusions and purchase authority. Use a short call only for ambiguity.','Buyer objective in their own words; complete inputs timestamp; decision-maker; qualify, rescope or decline reason.'),
('3 Feasible','Source owner checks supported predicates and source. Sibongile compares eligible set, distinct providers and capacity. Target answer within one working day of complete inputs.','Approved edition/hash; source exceptions; eligible provider/location set; required quantity; reviewer/production slot. A location count alone fails.'),
('4 Scope accepted','Sibongile drafts one version; Henry approves terms, fee and exceptions. Buyer accepts that exact version.','Geography, canonical service type, factual selection rule/window, exclusions, grouping, quantity, files, recipient, fee, terms, date and explicit acceptance.'),
('5 Paid','Authorised merchant operator issues the approved request after readiness and scope approval. Match authoritative merchant evidence; confirm by next working day.','Order/payment ID, paid status, amount, GBP, payer/order match, refund state. A screenshot or email claiming payment is insufficient.')
],[1.12,2.85,2.98])
h('No-supported-scope path')
p('If fewer than the agreed minimum qualify, stop before payment. Offer a supported wider scope or refuse. A smaller pack or different price is a separately approved commercial exception with revised terms and buyer acceptance; it is not an automatic fallback. Missing evidence stays missing. Never pad a shortlist or relabel locations as organisations.')
h('Payment exception path')
p('For mismatch, failed, duplicate or uncertain payment, open a case with a merchant owner and update date. Do not ask the buyer to pay again until the first attempt is resolved. Any refund or live merchant change needs the separate existing authorisation. Scope changes require a new version and acceptance; they cannot silently change what the paid buyer receives.')

page('Production to usable customer receipt')
table(['Stage','Owner and control','Exit evidence'],[
('6 Produced','Henry/source operator generates only through proven inputs; Sibongile curates and explains the frozen scope. Record actual work start, minutes and version.','Draft PDF, CSV and manifest; source/scope IDs; generation timestamp; exact inventory; defects visible.'),
('7 Reviewed','Separate competent reviewer checks every entry and rendered files. Producers fix defects; reviewer rechecks changed final version.','PASS with reviewer identity, date, scope version and SHA-256 for each final file. No PASS for a different draft.'),
('8 Released','Authorised delivery operator checks recipient, payment, entitlement, approved route and hash match. Send actual access/use instructions.','Approved final inventory, recipient/route, release timestamp and send/access receipt. An upload is not release.'),
('9 Access confirmed','Sibongile checks by next working day after release, handles bounced mail or link denial, and asks for explicit confirmation.','Agreed recipient confirms PDF, CSV and manifest are accessible and openable. Until then usable receipt remains pending.')
],[1.1,3.0,2.85])
h('Reviewer checklist for each pack')
p('Reconcile eligible set and shortlist count to the accepted predicates. Check provider/location grouping, every source URL, factual date, current-versus-change meaning, evidence edition and uncertainties. A location rating must not be stated as an organisation-wide rating. Check the fit explanation and qualification question without inventing procurement need.')
p('Open the actual PDF and CSV; reconcile counts/fields between them; check readability, import-safe CSV, empty values and unsupported filler. Verify attribution in the PDF and manifest, with CQC and the applicable Open Government Licence link. Confirm the recipient inventory contains the reviewed manifest as well as PDF and CSV.')
h('After receipt')
p('Record product acceptance separately: the buyer confirms the pack meets the accepted scope, or identifies a defect. Follow up on usefulness five working days after confirmed access. Ask which decision or action the pack supported and what remained missing. A compliment, download or silent buyer does not establish product acceptance or impact.')
p('Recovery remains owned until resolved. An access issue gets acknowledgement by next working day and a realistic next-action date. A factual defect freezes affected release/new commitments, receives an independent correction review, and follows the operative remedy. Do not invent a refund entitlement or guaranteed fix deadline in a template.')

page('Three clocks and the promised date')
p('Customer elapsed time runs from real enquiry receipt to confirmed usable access. Report calendar elapsed time and elapsed working days. CareGist waiting is measured at each owned queue. Labour is actual time worked by every role. Eight labour hours does not mean eight-hour delivery.')
pic('clocks.png','Figure 2. Illustrative cases starting Monday 26 October, conditional on readiness and capacity. Grey bands mark weekends; these are not forecasts.')
table(['Clock point','Prompt case','Late payment / later access'],[
('Enquiry / confirmation','Enquiry Mon 26 Oct; scope and source Tue 27 Oct.','Enquiry Mon 26 Oct; scope and source Thu 29 Oct.'),
('Original target / paid','Target Fri 30 Oct; verified paid Tue 27 Oct.','Target Tue 3 Nov; verified paid Fri 30 Oct.'),
('Date handling','Retain Fri 30 Oct if capacity supports it.','Buyer explicitly agrees Wed 4 Nov; original target stays recorded.'),
('Release / access','Fri 30 Oct release and access confirmation.','Wed 4 Nov release; Thu 5 Nov access confirmation.'),
('Elapsed / usefulness','4 calendar /4 elapsed working days; usefulness Fri 6 Nov.','10 calendar /8 elapsed working days; usefulness Thu 12 Nov, beyond pilot.')
],[1.25,2.75,2.95])
p('Published target: three English working days after the later of written scope confirmation and source-path check. Payment is required before work, but it does not start or automatically reset that clock. Keep the original target and explicitly re-agree a revised calendar date after late payment. Confirm operative terms and any accepted exception before work.')
p('If Friday release is not confirmed by the recipient until Monday, receipt is Monday. An acknowledgement task is a staff obligation, not a guaranteed buyer response. For a five-working-day usefulness interview by Tuesday 10 November, confirmed access must be no later than Tuesday 3 November; later access requires a funded follow-up beyond the pilot.')

page('Readiness before payment and review before release')
pic('gates.png','Figure 3. Readiness is established before charging; review of the paid order’s final files happens before sending.')
table(['Gate','Evidence to approve','Blocks'],[
('G0 Staffed','Roster/access, monitored inbox, one restricted register, backup handover and separate reviewer slots.','Paid delivery commitments.'),
('G1 Source / release','Exact deployed source identity; approved accessible inputs; edition/hash, freshness/completeness, material exceptions; applicable source, latency and shadow gates.','Unsupported feasibility, claims and affected source/export/commercial actions.'),
('G2 Scope fidelity','One supported contract and eligible set; distinct nonblank provider IDs; geography/service/category/window/exclusions/ranking consistent through output.','Guaranteed quantity or payment for an unproducible scope.'),
('G3 Outreach','Campaign/recipient permission, business type, provenance, relevance, privacy basis/information, suppression and approved copy; TPS/CTPS for calls; mail configuration checked.','Marketing sends/calls. Service communications require their own approved route.'),
('G4 Terms / money','Approved fee/tax/terms/date; selected merchant route; matching, failed/mismatched/duplicate payment, refund/cancellation and replay rehearsal.','Payment request and live money operations.'),
('G5 Full rehearsal','Full-sized synthetic pack, timed production, independent rendered-file review; real authorised controlled inbox/private access; wrong-recipient denial and recovery.','First paid commitment/payment request, not merely first paid release.'),
('G6 Each release','Accepted scope, matched paid status, every-row QA, final corrected hashes, attribution inventory and recipient/route authorisation.','Release of that specific order/version.')
],[1.08,4.05,1.82])
p('All operational approvals are unverified in this document. Record CLOSED, evidence-ready awaiting review, or approved with exact receipt/version/date. The date alone cannot pass a gate. If a required reviewer/provider is unavailable, that review remains uncompleted.')
p('Keep the project loop intact: Grok investigation and checkpoint → Codex named fix → Grok customer-path retest → DeepSeek frozen-evidence challenge → Hermes next bounded step. Model agreement does not approve a commercial release; this document’s independent review also does not substitute for that loop.')

page('Source and product work in the right order')
h('Current-status scope is a proposal to prove')
p('Local code separates Brief checkout flags from the Radar readiness helper. But the inspected Brief generator only includes rows with supported events in its time window. Removing change-based language does not turn it into a complete current-status selector. A current-status manual method needs approved selection/ranking, complete fields, lineage and full output review; it cannot be a gate bypass.')
p('Registration dates do not prove physical openings. “Not rated” has several possible meanings. Current ratings and dates must retain their published location/service level. Source edition age, effective fact date and time of observation are different. Weekly directory fields alone may not satisfy every requested fact; prove actual inputs before promising them. [S3]')
table(['Priority','Bounded task','Acceptance evidence'],[
('1','Pin deployed code and approved source; confirm dormant self-serve and delivery remain closed.','Exact release/input identities; applicable freshness/completeness/exception receipts. Local c17d1df is not proof of public 0fda3cac contents.'),
('2','Define one supported scope contract and eligibility count.','Three locations under one provider count as one. Reject blanks/malformed IDs; test category, service, region, exclusions and window. Compare complete eligible set before shortlist truncation.'),
('3','Prove actual production method and preflight.','Explicit generation date, approved input/hash; refuse stale, missing, conflicting or unsupported data. No raw CLI defaults for paid work. Reconcile PDF/CSV/manifest.'),
('4','Rehearse date, merchant and restricted delivery route.','Working-day date cases; payment mismatch/duplicate/replay; permissions, expiry or scheduled revocation, wrong-recipient denial, bounce/access recovery.'),
('5','Repair event labels and register mapping only where needed.','Clicks labelled intent; inbox defines receipt. CRM import maps contacts/deals/tasks, preserves suppression/history and reconciles accepted/rejected rows. Screening import stays separate.'),
('Separate','Investigate poll schedule and future automated commerce.','Actual rolling valid poll count ≥24, success ratio ≥90%, plus other applicable checks. More schedule slots are not a PASS; any merge/deploy is separately authorised.')
],[.85,2.73,3.37])
p('Engineering envelope proposed: up to 17h in the initial pilot, re-estimated after the first investigation. Do not assume the rough 55h backlog is approved, or spend time connecting endpoints before the actual pack/source/contract is proven. Henry’s production minutes are order labour; engineering work is logged separately so it is counted once.')

page('A daily schedule that fits five hours')
pic('capacity.png','Figure 4. The feedback’s follow-up workload adds 20 minutes beyond the daily limit. The proposed two modes keep the total at five hours; throughput must adapt.')
table(['UK window','Discovery day','Production day'],[
('09:00–09:30','Inbox, deadlines, source status and exceptions (0.5h).','Same priority check (0.5h).'),
('09:30–11:30','Prospect research, approved sends and due follow-ups (2h total).','09:30–10:30 prospecting/follow-ups (1h); 10:30–11:30 production (1h).'),
('11:30–12:00; 13:00–14:00','Discovery, source/rehearsal or scope work (1.5h).','Production/curation (1.5h).'),
('14:00–14:30','Review handoff, approved release or access check (0.5h).','Same (0.5h).'),
('14:30–15:00','Replies, register, time log and next-day plan (0.5h).','Same (0.5h).')
],[1.47,2.74,2.74])
p('On Fridays, the 14:00–15:00 review replaces the two existing half-hour blocks. Daily inbox checks still happen within that hour. Meetings, discovery calls, follow-ups and recovery consume scheduled time; they do not sit on top of it. When queues grow, protect existing paid commitments before new acquisition.')
h('Four weekly allocations')
table(['Week','Operations','Prospecting','Service / rehearsal','Handoffs / notes / review','Total'],[
('12–16 Oct','5h','5h','12h','3h','25h'),('19–23 Oct','5h','9h','8h','3h','25h'),('26–30 Oct','5h','7h','10h','3h','25h'),('2–6 Nov','5h','6h','10h','4h','25h')
],[1.25,1.0,1.14,1.3,1.31,.95])
p('These are operator allocations, including follow-ups. Henry, engineering and independent review are separate. No more than two paid orders active, and the second requires reserved production and reviewer slots. Optional 9–10 November work is an extra 10h only if explicitly funded; otherwise prepare the operator handover by 6 November and Henry owns the decision and remaining obligations.')

def week(title,intro,rows,gate,contingency):
    page(title,intro)
    table(['Date','Sibongile’s five-hour allocation','Henry / reviewer and end-of-day evidence'],rows,[.83,3.36,2.76])
    h('Friday decision');p(gate)
    h('If the week slips');p(contingency)

week('Week 1 establish and rehearse','12–16 October. Goal: a supported offer, named people and a full rehearsal. Outreach may begin only after its own permissions and controls pass; Monday has no mandatory send quota.',[
('Mon 12','1h roster/register/inbox; 1h tracker and source provenance; 1h scope draft; 1h candidate/source audit; 1h decision and next-action record.','Record all blocking decisions, loss ceiling and reviewer. Pin public/local evidence. Outcome: one supported proposed scope or explicit shortfall.'),
('Tue 13','1h operations and record; 1h permitted cohort verification; 2h eligibility/category/exclusions; 1h rehearsal preparation.','Count eligible distinct providers independently; test source method. Outcome: canonical scope and complete candidate set.'),
('Wed 14','1h operations; 1h cohort/sample work; 2.5h timed full-pack production; 0.5h review handoff.','Source owner supports generation; independent reviewer inspects draft. Outcome: full draft or honest incomplete record, actual minutes and defects.'),
('Thu 15','1h operations; 1h approved research/sends; 2.5h corrections and access/recovery rehearsal; 0.5h notes.','Test-mode merchant route, recipient denial and release evidence. Outcome: exact corrected files and route recovery receipts.'),
('Fri 16','1h operations; 1h approved research/sends; 2h re-review/rehearsal gaps; 1h weekly gate and dashboard review.','Independent reviewer rechecks final pack; Henry records remaining blocks. Outcome: G0–G5 evidence packet, each approval or explicit open item.')
], 'Do not call a three-entry sample a completed service rehearsal. Sample entries should be genuinely source-bound and independently reviewed; full rehearsal uses the complete agreed minimum shortlist and all delivered file types.',
'Reduce discovery volume and repeat the failed section of the rehearsal. Paid commitments wait. Do not open flags, invent a current-status route or assume source completeness to meet Friday.')

week('Week 2 discover and validate willingness','19–23 October. Goal: understand buyer decisions and reach a supportable readiness decision. A conversation, request for a scope and written scope acceptance are separate observations.',[
('Mon 19','1h operations; 2h approved research/sends/follow-ups; 1.5h warm-history review/discovery; 0.5h notes.','Verify prior correspondence and permission before re-approaching. Outcome: buyer decision/current method, not assumed purchase intent.'),
('Tue 20','1h operations; 2h approved prospecting/follow-ups; 1.5h discovery and feasibility; 0.5h notes.','Run supported eligibility checks only. Outcome: qualify/rescope/no decision and elapsed timestamps.'),
('Wed 21','1h operations; 2h approved prospecting/follow-ups; 1.5h rehearsal gaps or discovery; 0.5h notes.','Finish preflight/route defects. Outcome: corrected evidence packet and capacity estimate from measured work.'),
('Thu 22','1h operations; 2h approved prospecting/follow-ups; 1.5h one scoped opportunity; 0.5h notes.','Confirm terms, fee, source and reviewer slot. Outcome: versioned proposed scope if suitable; no payment yet.'),
('Fri 23','1h operations; 1h approved prospecting/follow-ups; 2h source/scope readiness; 1h weekly review.','Independent approvals for G0–G5 and campaign controls if complete. Outcome: readiness decision and first paid-slot reservation or continued closure.')
], 'Aim to learn from up to ten substantive conversations, without assuming ten are obtainable. A proposed learning trigger is at least three requests for a scope at £745 out of ten substantive conversations; it is a heuristic, not statistical validation.',
'If demand is weak, investigate the decision/problem and segment before scaling sends. If readiness fails, keep discovery within approved limits and fix the named gap. No calendar-driven invoice.')

week('Week 3 fulfil the first paid order','26–30 October. Earliest planning week for a paid order, conditional on approvals, buyer agreement and capacity. Zero orders remains a valid observed outcome.',[
('Mon 26','1h operations; 2h approved prospecting/follow-ups; 1.5h final suitability/scope; 0.5h notes.','Confirm G0–G5 and exact accepted scope; authorised invoice only after checks. Outcome: request issued or explicit pending reason.'),
('Tue 27','1h operations; 1h approved prospecting/follow-ups; 2.5h production if payment verified; 0.5h handoff.','Merchant match and real start timestamp; source/generator work logged as production. Outcome: draft inventory or unpaid/pending case.'),
('Wed 28','1h operations; 1h approved prospecting/follow-ups; 2.5h curation/correction; 0.5h handoff.','Independent reviewer checks every row and rendered files. Outcome: PASS or named defects, not automatic release.'),
('Thu 29','1h operations; 2h approved prospecting/follow-ups; 1.5h review/recovery or second feasibility; 0.5h notes.','Re-review corrected hashes and capacity. Outcome: G6 receipt or blocked release with buyer update.'),
('Fri 30','1h operations; 1h approved prospecting/follow-ups; 2h authorised release/access support; 1h weekly review.','Release only reviewed paid version to agreed recipient. Outcome: released, access pending or confirmed, and separate acceptance status.')
], 'A sample fast path agrees scope/source and receives payment on Tuesday 27 October, targeting Friday 30 October. The actual agreement fixes the date; this schedule is not a buyer promise. Book the +1 access check and +5 usefulness task from their actual anchors.',
'If payment is late or review fails, communicate before the agreed date and follow the operative remedy/date process. A second order is accepted only with a separate eligible scope and reserved production/review capacity.')

week('Week 4 prove value and close the pilot','2–10 November. Core operator work ends Friday 6 November. Protect already accepted obligations and distinguish missing feedback from failure.',[
('Mon 2','1h operations; 1.5h approved prospecting/follow-ups; 2h second-order production or recovery; 0.5h handoff.','Review current source and second-order capacity. Outcome: draft or no-order learning record; no assumed sale.'),
('Tue 3','1h operations; 1h approved prospecting/follow-ups; 2.5h review/release/access or corrections; 0.5h notes.','Latest same-day confirmed access for +5-working-day usefulness by Tue10. Outcome: G6 and actual access evidence, or funded follow-up tail.'),
('Wed 4','1h operations; 1.5h approved prospecting/follow-ups; 1.5h access/acceptance/usefulness; 1h labour and cost reconciliation.','Check paid, accepted and useful separately. Outcome: decision/action evidence, defects, complete labour ledger.'),
('Thu 5','1h operations; 1h approved prospecting/follow-ups; 2h customer follow-up/recovery; 1h final dashboard draft.','Third order only with proven capacity, source and independently reserved review. Outcome: explicit accept or decline decision.'),
('Fri 6','1h operations; 1h due follow-ups; 2h evidence/handover pack; 1h weekly review.','Reserve tail ownership and money; record all pending buyer/recovery tasks. Outcome: final operator handover and decision-ready evidence.')
], 'On Monday 9 November, Henry reconciles unresolved cases and prepares the decision. On Tuesday 10 November, record continue, limited extension, change or stop with reasons and next owner. Sibongile works these two days only if the extra 10h is explicitly approved; otherwise Henry/approved backup covers.',
'Later access shifts usefulness beyond Nov10. The decision is provisional while relevant feedback is pending. Never close open buyer obligations just because the pilot ends. Weekends have no planned operator hours; arrange cover if accepted obligations require it.')

page('Growth work without overloading the service')
h('Research priority and permission')
p('Start with the 20-row tracker: validate the 16 FIT decisions and resolve the four REVIEW rows. Seventeen have no recorded first contact; three have recorded contact. Fulcrum’s September sample request and soft deferral deserve a careful history check, not an assumption of consent or demand. No payment agreement is established. Review the actual correspondence before quoting it in a new message.')
p('Research relevant compliance/quality consultancies and then training suppliers in the chosen region. Record business identity, published contact source, fit, recipient type and suppression before any approved contact. Retained historical lists are research inputs. Do not import private contact details into Git or put them in the delivered plan.')
h('Volume is conditional')
p('Ten new first touches per day is an upper workload hypothesis, not a target that overrides five-hour capacity. Start with a small verified batch after G3 approval. Measure minutes per verified contact, reply load and deliverability. One approved follow-up after five working days is permitted only for a still-relevant non-replier who has not opted out; never treat it as mandatory.')
table(['Feedback volume scenario','Through 6 Nov','Tail obligation'],[
('10 new each core weekday','200 new first touches; up to 150 follow-ups.','Last week’s 50 follow-ups fall 9–13 Nov: 20 by Nov10 and 30 beyond it, if all remain eligible.'),
('Capacity-controlled final plan','Actual sends are limited by approved list and allocated hours. No fixed total forecast.','For each planned follow-up, reserve an owner/date within funded capacity. If tail is unfunded, reduce/stop new first touches or approve a shorter campaign cadence beforehand.')
],[2.05,2.4,2.5])
h('Campaign safeguards')
p('For this pilot, restrict new cold email research to verified corporate subscribers such as companies/LLPs. A named business address may still contain personal data. Record the applicable privacy basis, information and objections; corporate status alone is not a complete permission record. Individual subscriber rules differ; do not assume a cold soft opt-in. Screen marketing calls against TPS, CTPS and prior objections. [S2]')
p('Approve actual copy, sender identity/privacy link and mail setup, including SPF/DKIM/DMARC checks. Log individual sends and replies. Investigate bounces; pause affected sending for a spam complaint, suppression breach or a deliverability pattern. Restart only after the cause is resolved and the relevant control is re-approved, not automatically the next morning.')
p('Discovery asks: what decision are you making, how do you prioritise now, what evidence would change your decision, what would the sample miss, and would you request a written scope at £745? Record objections and counterexamples, not only positive comments.')

page('Cost model and loss ceiling','All figures below are illustrative planning assumptions, not observed service cost or merchant pricing. Use actual invoices, hours and refunds in the decision.')
pic('economics.png','Figure 5. £745 price; illustrative £30/hour and £25 other cost/order. Whole-pilot model includes 100 operator hours, 17 engineering hours, 3h rehearsal review and 1.5h separate review/order; founder time is additional.')
table(['Model','Calculation','Meaning'],[
('Unmeasured hybrid','5h × £30 + £25 =£175; contribution £570 (76.5%).','A hypothesis only. If review rises from 1h to 1.5h, other work unchanged, total becomes 5.5h and contribution £555.'),
('Original manual','8h × £30 + £25 =£265; contribution £480 (64.4%).','Includes 1h allocated acquisition and 1h review; not measured elapsed delivery time.'),
('Manual with 1.5h review','8.5h × £30 + £25 =£280; contribution £465 (62.4%).','Same underlying assumptions, with an extra 0.5h independent review.'),
('60% contribution floor','(£745 × 40% − £25) / £30 =9.1h.','Only under this price/cost model. Unincluded fees/tax/refunds or a different net price reduce the available hours.'),
('Two-order whole pilot','£1,490 − £3,000 operator − £510 engineering − £90 rehearsal review − £90 order review − £50 other =−£2,250.','Before founder/backup hours, extra review/recovery costs and any tax adjustment. Do not subtract order operator labour again.')
],[1.43,2.77,2.75])
p('At zero, one and three paid orders the same fully used envelope yields −£3,600, −£2,925 and −£1,575 respectively, before the exclusions above. Rehearsal review reserves two 1.5h slots provisionally; all additional review is logged and costed. These are scenarios, not probabilities. Optional 10 operator hours adds £300. Use actual hours/costs; the envelope is not a spending instruction.')
p('The original 55h build at £30/h costs £1,650. A hypothetical three-hour saving (£90/order) needs 19 whole Briefs to recover it; an initial 17h (£510) needs six. Saving and scope are unproven. Henry must set the actual cash ceiling, separate labour allocations, loss tolerance and escalation point explicitly; a profitable-looking marginal order does not establish a profitable pilot.')

page('One register and the 10 November decision')
h('Minimum restricted record')
p('Use one authoritative operational store. Prospect fields: ID, organisation, contact source, business/recipient type, relevance, privacy/suppression evidence, contact history, send/template dates, follow-up eligibility/due date and outcome. Enquiry fields: received/acknowledged/complete-input timestamps, owner, objective, next action/date and blocker. Order fields: scope version, eligible set/count, source identity/checksum, terms/fee/date/recipient, approvals and merchant match.')
p('Record work, review and release as separate dimensions: work started/prepared; review pending/PASS/defects; released/access pending/confirmed; product acceptance pending/accepted/defect; remedy open/resolved. Store exact file hashes, reviewer receipt and actual minutes by role. Preserve originals and history; do not overwrite a rejected scope with a later one.')
table(['Dashboard item','Count only this evidence'],[
('Research / actual sends','Verified prospect rows / actual authorised sent messages; drafts are separate.'),
('Substantive conversations','Buyer described a decision and current method. Link one buyer across repeated replies.'),
('Requests / accepted scopes','Explicit request for scope at £745 / explicit written acceptance of a specific version.'),
('Payment / release / access','Authoritative matched paid status / reviewed authorised release / recipient confirms all files open.'),
('Acceptance / useful decision','Buyer confirms scope met / names a decision or action supported; keep distinct.'),
('Time / money / problems','All labour, including unsuccessful acquisition; gross receipts, fees, refunds and net separately; pending, overdue and recovery cases with owners.')
],[2.08,4.87])
h('Decision rule')
p('Continue only if the service is safe, two paid Briefs are explicitly accepted, two buyers describe a supported decision/action, full costs are understood and capacity fits the approved budget. The illustrative 9.1h ceiling is a planning check, not a universal business threshold. Keep the small cohort as case studies, not a statistically proven conversion rate.')
p('Extend narrowly if readiness is proven but demand or buyer feedback is still pending, with a new owner/date and funded ceiling. Change the segment/offer question after weak substantive demand; do not scale the same unvalidated message. Stop new commitments if source/review controls fail, a material defect is unresolved or spending reaches the approved limit. Continue recovery duties and contract remedies for existing buyers.')
p('For conversions, use linked mature cohorts and show numerator, denominator and pending cases. Actual median/p90 journey time, abandonment, satisfaction and conversion baselines are unknown today; do not manufacture a funnel graph from the planning target.')

page('Draft messages for enquiry and scope','These are reviewable drafts. Use only after the relevant copy, staffing, terms and communication route are approved. Replace every bracket and remove unsupported options. Use verified trading identity and privacy information; no unverified company details are supplied here.')
h('T1 and T5 acknowledgement')
p('Thank you for your message. Your reference is [ID], and [name] will look after your enquiry. To check the scope, we need [only missing items]. We will update you by [date] with the next step. [For an approved auto-reply, state only the tested staffed response promise; a reference is issued after the real record exists.]')
h('Qualification prompts')
p('What do you sell and to which provider types? Which territory decision are you making, and by when? How do you choose accounts today? What providers or areas should be excluded? Who approves a purchase? What would make these three reviewed sample entries useful or irrelevant? A 15–20-minute call is optional when written answers leave ambiguity.')
h('T6 scope acceptance')
p('Reference [ID], scope version [v]. Your decision: [objective]. Territory: [canonical area]. Provider types: [taxonomy]. Selection rule and evidence window: [precise supported predicate, status/change definition and dates]. Exclusions/grouping: [rules]. Shortlist: [n] distinct CQC provider IDs, with relevant location facts identified. Source: [approved edition, checked date and limitations]. Files: executive PDF, CSV and source/licence manifest. Fee and tax treatment: [approved £745 terms]. Recipient: [agreed recipient]. Operative terms: [version/link].')
p('Written scope confirmation: [date]. Source-path check: [date]. Delivery target: [calendar date], three English working days after the later of those confirmations, under the operative terms. Payment is required before work; if payment timing affects capacity, we will agree a revised date explicitly. Excludes contact enrichment, forecasts, future inspections, compliance advice and exclusivity. Please confirm acceptance of scope version [v].')
h('T7 payment and T8 start confirmation')
p('After acceptance and readiness: Your approved [merchant route] payment request is for [£745 GBP] and reference [ID], scope [v]. Work starts after payment is verified. Please do not pay twice; reply here if the payment status looks unclear. [Give only the tested route and actual date.]')
p('After verified payment: We have matched your payment to [ID]. [Work has started on date / work is scheduled to start on date]. Your agreed delivery date is [date], and the next update is [date]. We will tell you promptly if that date is at risk. Do not state “started” before work actually starts.')

page('Draft messages for release and growth')
h('T9 release and T10 access check')
p('Your Territory Brief [ID] is available through [approved route]. It contains the executive PDF, CSV and source/licence manifest, based on [edition]. Access is for [agreed recipient]; [describe only tested expiry/revocation settings and any recipient sign-in requirement]. [One short instruction for opening/importing the CSV.] Please confirm that all three files are accessible and openable. For support, reply with [ID].')
p('Access check: Can you confirm that the PDF, CSV and manifest all open? If something fails, please tell us which file and what happens. We will respond by [realistic date] with the next step. This is a response date, not an invented guarantee that every issue will be fixed by then.')
h('T11 late payment and T12 insufficient scope')
p('The original scope/source confirmations for [ID] were [dates], with delivery targeted for [original date]. Payment was verified on [date]. We propose [new date] because [specific capacity impact]. Please confirm agreement; we will record the revised date before starting. Until then, the date change remains unresolved under the operative terms.')
p('Our checked source supports [n] distinct eligible providers for [criteria], fewer than the agreed minimum. No payment has been taken. We can discuss [supported wider scope] or stop. A smaller pack or different price would require a separately approved written agreement. We will not pad the shortlist.')
h('T2 warm contact and T13 previous contact')
p('Warm contact, only after correspondence verification: You asked about a sample in [month]. We now propose a one-off evidence-linked Territory Brief for an agreed account decision. Is that decision still relevant? If so, we can show three reviewed sample entries before discussing a scope. If you would rather not hear from us, please say so and we will suppress further marketing.')
p('Previous contact: I wrote in [month] about [actual previous offer]. Our proposed pilot is now a one-off Territory Brief: an agreed shortlist with the source facts and qualification questions behind each entry. How does your team currently choose providers to investigate? [Approved opt-out and privacy information.] Do not assert that the new offer is “more useful” before the buyer establishes that.')
h('T3 first touch and T4 one follow-up')
p('First touch: I am contacting you because [verified organisation] supplies [service] to adult social care providers. CareGist proposes evidence-linked research for an agreed territory and account criteria. How does your team choose which providers to investigate first? Would a short conversation about that decision be useful? [Approved identity, opt-out and privacy information.]')
p('Follow-up: One follow-up on my note about provider research in [region]. If it is not relevant, please let me know and I will not send further marketing. [Use only for an eligible non-replier after the approved interval; do not follow up after an objection or substantive reply.]')

page('Sources and evidence boundaries')
p('Prepared 11 October 2026. Public snapshots below were captured earlier today; they are dated service assertions, not independent reprocessing of all source records. No real enquiry, payment, refund, campaign or buyer delivery was executed for this document.')
table(['ID','Source and use'],[
('S1','User-provided Claude Operating Bible: https://claude.ai/artifact/PvG6FeFCZ4aFkSo5gPEqoT#56007739-7555. Full visible artifact reviewed through the browser. Original Hybrid Operating Model: https://claude.ai/artifact/9yTh43J4rZrZ8weELDqSU7#092a092b-83bc. Both are proposals, not receipts.'),
('S2','ICO business-to-business marketing: https://ico.org.uk/for-organisations/direct-marketing-and-privacy-and-electronic-communications/business-to-business-marketing/ . Primary guidance checked for corporate/individual subscribers, personal data and objections. Actual campaign review remains required.'),
('S3','CQC using CQC data: https://www.cqc.org.uk/about-us/transparency/using-cqc-data . Primary guidance checked for fields, rating levels, limitations and attribution. Applicable licence: https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/ .'),
('S4','Public CareGist Territory page and terms: https://caregist.co.uk/pricing/territory and https://caregist.co.uk/terms . Saved in artifacts/customer-journey/2026-10-11/{territory,terms,health}-public.json. At 08:01 UTC, reported release 0fda3cac…; checkoutReady, shadowCoveragePassed and deliveryEnabled false.'),
('S5','Earlier source-grounded work: artifacts/blueprint-30-days/2026-10-11/BLUEPRINT.md and REVIEW.md; artifacts/customer-journey/2026-10-11/ENQUIRY_TO_RECEIPT.md and PRINCIPAL_GROWTH_PM.md; artifacts/strategy-review/2026-10-11/{PLAN,MANUAL_RUNBOOK,OPERATOR_RECORD,EVIDENCE}.md. User pasted red-team assessment dated 11 October.'),
('S6','Local checkout c17d1df9b6214664d650274ac12e9f181fcf5ad7: api/routers/billing.py; api/services/territory_brief.py and territory_brief_fulfilment.py; tools/generate_territory_opportunity_brief.py; frontend/lib/directory-db.ts; leads/README.md. Local inspection does not establish deployed contents.'),
('S7','VA_REVENUE_TRACKER.csv: 20 rows, count/status and recorded warm history inspected. CareGist_VA_Sales_Training_Guide_v2.docx: South Africa time-zone guidance inspected. Neither proves current permission, availability or actual pipeline demand.')
],[.65,6.3])
h('What remains to be proven operationally')
p('Real staff account access and inbox handling; current approved source completeness and supported production path; prescribed independent commercial challenge; live merchant configuration/settlement and permissions; private recipient-only access and recovery; measured end-to-end labour; buyer acceptance and usefulness. Exact deployed source is not available in the inspected local checkout. This final plan preserves those limits.')
p('Graphs show a process or explicitly labelled illustrative assumptions. There is no measured sales funnel or conversion uplift. Document completion and commercial readiness are separate decisions. Adopt this proposed schedule through explicit decisions and actual gate receipts; do not use its final formatting as evidence that a sale is approved.')

doc.save(OUT/'CareGist_30_Day_Operating_Plan.docx')
(OUT/'FINAL_PLAN.md').write_text('\n'.join(source),encoding='utf-8')
print(json.dumps({'docx':str(OUT/'CareGist_30_Day_Operating_Plan.docx'),'sections':page_num,'words':len(' '.join(source).split())}))
