from pathlib import Path
import hashlib, json, shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle, FancyArrowPatch

OUT=Path(__file__).parent
FIG=OUT/'figures'; FIG.mkdir(parents=True,exist_ok=True)
NAVY='#203B50'; TEAL='#24756F'; GREY='#58636D'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':12})

fig,ax=plt.subplots(figsize=(10,4.1)); ax.axis('off'); ax.set(xlim=(0,10),ylim=(0,4.2))
steps=[('1 Receive','Read the actual enquiry'),('2 Ask','Understand the buyer'),('3 Check','Can we deliver this?'),('4 Agree','Write down the exact job'),('5 Verify','Check the payment'),('6 Make','Prepare the three files'),('7 Review','A separate person checks'),('8 Release','Use the approved route'),('9 Confirm','Buyer opens all three files')]
positions=[(0,3),(3.5,3),(7,3),(7,1.65),(3.5,1.65),(0,1.65),(0,.3),(3.5,.3),(7,.3)]
for i,((title,desc),(x,y)) in enumerate(zip(steps,positions)):
    ax.add_patch(Rectangle((x,y),3,.9,facecolor='#F1F5F6',edgecolor='#B5C1CA',lw=.8))
    ax.text(x+.13,y+.58,title,fontsize=14,fontweight='bold',color=NAVY)
    ax.text(x+.13,y+.22,desc,fontsize=10.5,color=GREY)
    if i<8:
        nx,ny=positions[i+1]
        if y==ny: a=(x+3,y+.45) if nx>x else (x,y+.45); b=(nx,ny+.45) if nx>x else (nx+3,ny+.45)
        else: a=(x+1.5,y);b=(nx+1.5,ny+.9)
        ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=13,color=TEAL,lw=1.5))
fig.savefig(FIG/'journey.png',dpi=220,bbox_inches='tight',facecolor='white');plt.close(fig)

fig,ax=plt.subplots(figsize=(10,2.6)); ax.set(xlim=(0,5),ylim=(-.6,1.6));ax.axis('off')
labels=['Tue 27 Oct','Wed 28 Oct','Thu 29 Oct','Fri 30 Oct']
notes=['Scope + source\nchecked','Working day 1','Working day 2','Working day 3\nTarget release']
for i,(label,note) in enumerate(zip(labels,notes)):
    x=.4+1.35*i; ax.scatter(x,.9,s=100,color=TEAL);ax.text(x,1.25,label,ha='center',fontsize=13,fontweight='bold',color=NAVY);ax.text(x,.35,note,ha='center',va='top',fontsize=11.5,color=GREY)
    if i<3:ax.add_patch(FancyArrowPatch((x+.15,.9),(x+1.2,.9),arrowstyle='-|>',color=TEAL,mutation_scale=12,lw=1.5))
fig.savefig(FIG/'clock.png',dpi=220,bbox_inches='tight',facecolor='white');plt.close(fig)

fig,ax=plt.subplots(figsize=(10,2.5));ax.axis('off');ax.set(xlim=(0,5.1),ylim=(-.1,1.7))
parts=[(.5,'Open and check'),(1,'Prospect'),(1,'Make the pack'),(1.5,'Serve customers'),(.5,'Handover'),(.5,'Replies and log')]
colors=[GREY,'#A87821',TEAL,'#397BB1','#725A8C','#687781'];left=0
for (hours,label),color in zip(parts,colors):
    ax.add_patch(Rectangle((left,.9),hours,.45,facecolor=color));ax.text(left+hours/2,1.125,f'{hours:g}h',ha='center',va='center',color='white',fontsize=12,fontweight='bold');left+=hours
for i,((hours,label),color) in enumerate(zip(parts,colors)):
    col=i%3;row=i//3;ax.add_patch(Rectangle((col*1.75,.47-row*.4),.1,.1,facecolor=color));ax.text(col*1.75+.16,.52-row*.4,label,fontsize=10.3,va='center',color=NAVY)
ax.text(0,1.55,'One production day = five working hours',fontsize=14,fontweight='bold',color=NAVY)
fig.savefig(FIG/'day.png',dpi=220,bbox_inches='tight',facecolor='white');plt.close(fig)

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.28)
for name in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    if s.element.pPr is not None:
        for el in list(s.element.pPr):
            if el.tag==qn('w:pBdr'):s.element.pPr.remove(el)
doc.styles['Normal'].font.size=Pt(12)
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',27),('Heading 1',21),('Heading 2',14)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].paragraph_format.space_before=Pt(10);doc.styles[name].paragraph_format.space_after=Pt(8)
doc.styles['Title'].paragraph_format.space_before=Pt(0)
hp=sec.header.paragraphs[0];hp.text='CAREGIST  |  SIBONGILE AND HENRY  |  TRAINING';hp.runs[0].font.size=Pt(8)
fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('11 October 2026  |  Practice guide  |  ').font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');fp._p.append(field)
doc.core_properties.title='CareGist training for Sibongile and Henry'
doc.core_properties.subject='Plain language practice workbook from enquiry to confirmed customer receipt'
doc.core_properties.author='CareGist training preparation'
source=[];page_count=0
def p(text,bold=False):
    pp=doc.add_paragraph();pp.add_run(text).bold=bold;source.append(text+'\n');return pp
def page(title,intro=None):
    global page_count
    page_count+=1;pp=doc.add_heading(title,0 if page_count==1 else 1);pp.paragraph_format.page_break_before=page_count>1
    source.append('\n# '+title+'\n')
    if intro:p(intro)
def h(text):doc.add_heading(text,2);source.append('\n## '+text+'\n')
def pic(name,caption):
    pp=doc.add_paragraph();r=pp.add_run();r.add_picture(str(FIG/name),width=Inches(7))
    inline=r._r.xpath('.//wp:docPr')[0];inline.set('descr',caption)
    pp=doc.add_paragraph(caption);pp.runs[0].font.size=Pt(10);source.append('[Figure: '+caption+']\n')
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w);t.rows[0].cells[i].text=headers[i]
    for row in rows:
        cells=t.add_row().cells
        for i,val in enumerate(row):cells[i].text=str(val)
    for ri,row in enumerate(t.rows):
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr()
            bd=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bd.append(el)
            pr.append(bd);mar=OxmlElement('w:tcMar')
            for edge,val in [('top','95'),('bottom','95'),('left','110'),('right','110')]:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),val);el.set(qn('w:type'),'dxa');mar.append(el)
            pr.append(mar);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'203B50' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.line_spacing=1.04
                for r in pp.runs:r.font.size=Pt(11);r.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    source.append('| '+' | '.join(headers)+' |\n')
    for row in rows:source.append('| '+' | '.join(str(x) for x in row)+' |\n')
def question(text,lines=1):
    p(text)
    for _ in range(lines):p('________________________________________________________________')

page('CareGist training for Sibongile and Henry')
p('A simple practice workbook from the first enquiry to the customer opening the product.',True)
p('Prepared 11 October 2026. Read this together. Use short sentences when you explain the service. Ask each other to show the evidence for each step.')
h('What you should be able to do')
p('Explain what we sell. Know who does each job. Keep the customer informed. Stop when something is missing. Hand over an order that another person can understand.')
pic('journey.png','The nine steps. An order reaches receipt when the agreed customer opens all three files and confirms this.')
h('How to use this workbook')
p('First, learn the offer and your role on pages 2 to 5. Then follow the order on pages 6 to 9. Practise on page 11 with the cards on pages 18 and 19. Use the handover sheet on page 12 and daily cards on page 15.')
p('This is practice material. It does not open checkout, approve marketing, take payment or allow delivery. Henry must record the required approvals. A separate reviewer must check a pack they did not make.',True)

page('What we sell','We help a supplier decide which care organisations to investigate for possible business.')
h('Who is the customer')
p('The customer is a supplier. For example, they may sell staff training or quality support to care organisations. The care organisations in our list are possible accounts for that supplier. They are not the person buying our report.')
h('What the customer buys')
p('The Territory Opportunity Brief costs £745 once. Henry must confirm how tax is handled before anyone quotes the total. It is a research pack for one agreed area in England.')
p('Our starting aim is 25 different provider IDs. A provider ID identifies a registered organisation. We must prove that 25 fit the agreed rules before asking for payment. The published range is 25 to 50. More rows do not always mean more organisations.')
table(['File','What it does'],[
('PDF report','Explains the findings in an easy-to-read report.'),
('CSV list','A list that spreadsheet software can open. It is not a special Excel workbook. Test any customer-system import before promising it.'),
('Source and licence manifest','A record of where the facts came from and the rules for reusing them.')],[2.15,4.85])
h('What we can say')
p('We show facts from named sources and explain why a provider may fit the agreed research rules. We suggest questions the supplier can ask next.')
p('We cannot promise that a provider needs training, has money to spend or will buy. A rating or registration date does not prove those things. We do not promise sales or personal contact details.')
p('Say: “This pack helps you decide where to investigate. You still need to speak to each organisation and check its needs.”')

page('A few words you need','Use these simple meanings during training. Say the full meaning when a customer may not know the word.')
table(['Word','Simple meaning'],[
('Enquiry','An actual message or call asking about the service.'),
('Scope','The exact job agreed in writing: what is included, left out and due.'),
('CQC','Care Quality Commission. It regulates health and adult social care in England.'),
('Provider','A registered organisation. It may have several locations.'),
('Location','A registered care site or service. It is not automatically a separate provider.'),
('Evidence','A record that shows what happened or supports a fact.'),
('Gate','A required check. If it has not passed, stop at that step.'),
('Independent reviewer','A competent person who checks a pack they did not help make.'),
('File fingerprint or hash','A code for the exact file version. If the file changes, recheck it.'),
('Do-not-contact list','The record of people or organisations that must not receive marketing.'),
('Receipt','The agreed recipient confirms they can open all three files.'),
('Acceptance','The buyer confirms that the delivered pack meets the agreed scope.')],[2.05,4.95])
h('Small counting example')
p('These are made-up training labels, not real CQC records. Training Provider A has two locations. Training Provider B has one. Training Provider C has one. That is 4 locations but only 3 providers.')
p('A location rating belongs at its stated level. One branch rating does not describe every branch or the whole organisation. A registration date is not proof that a new building just opened.')

page('Sibongile owns the customer journey','Your job is to keep each enquiry clear, recorded and moving through approved steps.')
h('When an enquiry arrives')
p('1. Read the actual message. Opening an email window on the website is not an enquiry received. Put the received time and a reference in the approved private register.')
p('2. Record what the buyer sells and what decision they want to make. Ask only for missing information. Name the next action, its owner and the next update date.')
p('3. Send an acknowledgement through the approved route only when Henry has explicitly authorised it. The proposed target is the next working day. Do not promise a date you cannot meet.')
h('Before the buyer pays')
p('4. Ask Henry or the source owner to confirm that the source, selection rules, quantity and timetable work. Check that a separate reviewer has a slot.')
p('5. Draft the scope. Henry checks the commercial details. The buyer must accept the exact version. Keep the acceptance with the order.')
h('When a paid order is ready to make')
p('6. Use only the approved instructions and source. Keep facts, source links, provider IDs, location IDs and doubts clear. Log the time you use.')
p('7. Hand the exact PDF, CSV and source record to the separate reviewer. Fix errors. Any changed final file needs another check.')
p('8. After release is authorised, use the approved recipient and delivery route. Ask the buyer to open all three files. Record their reply. Then check acceptance and later usefulness.')
h('Your stopping rule')
p('If payment, source, scope, review, permission or access is unclear, stop that step. Tell Henry what is missing. Give the customer a realistic next update date through the authorised route.')
p('You cannot independently approve a pack you helped make. Do not offer a new fee, refund or deadline exception on your own.',True)

page('Henry owns the operating checks','Your job is to make sure we can keep the promise before anyone commits money or releases files.')
h('Before any paid commitment')
p('Name the people, backup and separate reviewer. Confirm their access and available time. Approve the supported offer, terms, tax treatment and written delivery date.')
p('Prove the chosen source and method. Record the source version, date collected and file fingerprint. Check that it is fresh and complete enough for this scope. Resolve important gaps.')
p('A list of changes during a date window is not automatically a complete list of providers active today. Only promise the view the approved method can actually produce. Do not bypass source checks or change gate settings to make a job pass.')
p('Prove the payment and delivery routes with a full practice pack and controlled access test. Keep the real evidence. The short role-play in this workbook is not that full test.')
h('For each real order')
p('Check the record in our approved payment system. Match its paid status, amount, GBP currency (British pounds), payer, order reference and refund state. A buyer screenshot or a “paid” email is not enough.')
p('Before release, check the accepted scope, verified payment, separate review of the exact final files, correct recipient and approved route. Record explicit release authority. Missing permission stays missing.')
h('When something goes wrong')
p('Own payment problems, source problems and deadline decisions. Resolve an uncertain or duplicate payment before asking the buyer to pay again. Refunds and live merchant changes need their specific authorisation.')
p('If you generate, edit or curate a pack, you are a producer. You cannot be its independent reviewer. Find a separate competent person before accepting the paid commitment.',True)
p('You also cover the inbox only when your backup slot and access are confirmed. If nobody can cover a deadline, pause new commitments and arrange the next update for existing customers.')

page('From enquiry to written agreement','Do the checks in this order. Make one clear record for the order.')
h('Ask six useful questions')
p('What do you sell? Which care providers do you want to investigate? Which area? What decision will this pack help you make? When do you need it? Who can agree the purchase?')
p('Then ask about exclusions and provider types if these are still missing. For example, “Should we leave out providers you already work with?” Keep private customer information in the approved restricted system.')
h('Check whether we can do the job')
p('Can the approved source support the requested facts and selection rules? Are there enough different providers? Can the team and reviewer meet the date? If any answer is uncertain, do not ask for payment.')
p('If only 22 providers qualify for a 25-provider job, stop. Offer a supported wider scope for approval, or decline. A smaller pack or changed price needs Henry’s approval and the buyer’s agreement. Never add unsuitable entries to reach 25.')
h('Write the scope so both sides agree')
p('Include the area; provider types; factual selection rule and any date window; exclusions; how providers and groups are counted; quantity; source and source-check date; the three files; recipient; £745 price and tax treatment; terms; and a calendar delivery date.')
p('Save the version the buyer accepts. A change needs a new agreed version. Do not quietly change the job after payment.')
h('A useful customer update')
p('Draft only, for Henry to authorise before sending: “Thanks for your enquiry. I am checking whether we can support [area and rules]. I will update you by [date]. We will agree the exact scope and delivery date before asking for payment.”')
p('Every record needs four things: what happened, what happens next, who owns it and the next date.',True)

page('How the delivery clock works','The documented delivery target is three English working days. A working day excludes weekends and English public holidays.')
h('Find the later of two checks')
p('Start from the later of: the buyer confirming the written scope, and the source path being checked. A source path check confirms that our approved way of getting the needed facts works. Count the next working day as day 1. Agree a calendar date with the buyer.')
pic('clock.png','Fictional example in 2026. Both checks finish on Tuesday 27 October. The three-working-day target is Friday 30 October. Payment is verified before production.')
p('If the scope is confirmed on Tuesday 27 October but the source check finishes on Wednesday 28 October, the target becomes Monday 2 November. The weekend is skipped.')
h('What payment changes')
p('Real production starts only after matched payment. Payment does not restart the clock. If payment arrives too late for the agreed date, Henry must arrange a new written calendar date with the buyer before production begins.')
h('Keep four dates separate')
table(['Date','What it proves'],[
('Release date','We released the approved pack through the approved route.'),
('Access confirmation date','The agreed recipient said all three files open.'),
('Acceptance date','The buyer said the pack meets the scope.'),
('Usefulness check date','We asked what decision or action the pack helped with.')],[2.25,4.75])
p('Proposed targets: check access on the next working day after release; ask about usefulness five working days after confirmed access. If access is Friday 30 October, the usefulness check is Friday 6 November. These become promises only after approval and staffing.')

page('Make review and release the pack','Use the accepted scope as the checklist. A finished-looking file can still be wrong.')
h('Make the three files match')
p('The PDF, CSV and source record must describe the same shortlist. Check different provider IDs, relevant location IDs, factual reasons, sources and visible doubts. Leave out unsupported claims about needs, money or future events.')
p('Credit CQC where its data is used and record the applicable reuse terms. Check important source delays or missing data. Do not guess to fill a gap.')
h('Give it to the separate reviewer')
p('The reviewer checks every shortlisted entry and the actual PDF, CSV and source record against the accepted scope. They record PASS or the errors to fix. PASS means the checked version meets the required review; it does not replace release authority.')
p('Keep the file fingerprints with the review. After any edit, ask for the changed final files to be checked again. Neither of you can independently approve a pack you helped produce.')
h('Release only after the order check')
p('Henry checks the accepted scope, matched payment, reviewer result, exact file versions, recipient and approved delivery route. Record his explicit authority before the authorised operator releases the pack.')
p('Use the tested private route. Do not promise link expiry unless the actual system supports and has tested it. Confirm the intended recipient; a public link or wrong recipient is not acceptable.')
h('Confirm that the buyer has the product')
p('Draft only, for approval: “Please confirm that you can open the PDF report, CSV list and source record. If a file will not open, tell me which one and what happens.”')
p('Sending a message, seeing a download or getting no reply does not prove receipt. Keep the buyer’s clear confirmation. Then ask whether the pack meets the scope. Ask about its usefulness later.')

page('When to stop and ask Henry','Stop the affected step, record the problem, name the owner and give a realistic next update date through the authorised route.')
table(['Problem','What to do next'],[
('Source is old, incomplete or cannot support the request','Pause that job. Ask the source owner to resolve it or propose an approved scope change.'),
('Only 22 different providers qualify','Pause before payment. Rescope with approval or decline. Do not pad the list.'),
('Payment is unclear or counted twice','Pause production and release. Henry resolves the merchant record. Do not ask for another payment yet.'),
('You helped make the pack and are asked to approve it','Ask for a separate competent reviewer. Keep the pack waiting.'),
('Buyer cannot open a file','Record which file and the error. Test the authorised recovery route, arrange a next update and get fresh access confirmation.'),
('Wrong person receives the pack','Stop access where the approved route allows it. Tell Henry immediately and follow the incident process.'),
('Customer wants a refund or changed deadline','Henry checks the terms and authority. Record the agreed decision; do not promise it yourself.'),
('Someone asks to stop marketing','Record the request in the do-not-contact system immediately. Stop marketing to that recipient. Report any missed request or complaint to Henry.'),
('The work will exceed five hours today','Tell Henry early. Reduce new prospecting or rearrange approved work. Do not assume extra hours.')],[2.75,4.25])
p('A prospect’s public email address is not automatic permission to market. Henry must approve the campaign, recipients and applicable privacy checks. Calling also needs the required screening checks. Practise messages here without sending them.')

page('A four hour training session','Use two short learning sessions, then a practice order and a check. This is a proposed agenda, not proof that anyone has completed training.')
table(['Part','Time','Do this'],[
('Session 1','90 minutes','10: explain the offer. 20: count providers and learn the words. 25: practise an enquiry and scope. 20: explain each role. 15: work out dates.'),
('Session 2','90 minutes','20: Henry explains the operating checks. 25: use the fictional file and review cards. 25: practise access failure and payment problems. 20: use the handover sheet and daily cards.'),
('Joint role-play','45 minutes','Use the fictional order on page 11. Swap buyer and operator roles. Stop at each missing check.'),
('Learning check','15 minutes','Answer page 13 without the answer key. Explain any wrong answer, then practise that step again.')],[1.4,1.1,4.5])
h('Keep the training inside the time budget')
p('This takes four hours for each participant. Sibongile’s four hours come out of the planned 25-hour week; reduce other work to make room. Henry must reserve his own time. A full pack rehearsal and a separate reviewer need additional booked time.')
h('Suggested first week')
p('Monday 12 October: Henry records owners, approvals needed and access gaps. Tuesday 13 October: session 1. Wednesday 14 October: session 2. Thursday 15 October: role-play and learning check. Friday 16 October: discuss mistakes and book any repeat practice.')
p('These are proposed dates. If a slot has passed or access is missing, rebook it. Do not turn a planned date into a completed task.')
h('How the trainer should teach')
p('Explain one step. Show an example. Ask the learner to do it. Ask them to explain why. Let them use the workbook during practice. Keep the tone respectful; simple language is for clarity.')
p('The trainer can use the answers on page 14. Passing the learning check does not pass the operating gates or authorise real messages, payments or delivery.',True)

page('Practise one fictional order','Everything on this page is made up for training. Do not contact anyone, send a live payment request or release customer files.')
h('The buyer')
p('Pat works for Practice Training Company. Pat wants to decide which adult social care provider organisations to investigate for training work in a chosen English area. Pat asks for a 25-provider pack. The exact provider types, factual rules and exclusions still need agreement.')
h('Round one in 15 minutes')
p('Sibongile reads the pretend enquiry, gives it a practice reference and asks the missing questions. Henry plays Pat. Write a draft scope and an update date. Henry then explains what evidence is needed before the order can move on.')
p('Trainer gives this fault: “The source supports only 22 different providers.” Stop before payment. Explain an approved wider scope or a decline. Do not claim the fault is fixed just because you discussed it.')
h('Round two in 15 minutes')
p('For this part only, use the cards on pages 18 and 19. They describe an agreed 25-provider scope, source check, readiness approvals, simulated payment, three pretend files and a separate pretend reviewer result. They are paper practice props, not real gate evidence.')
p('Assume scope, source check and simulated payment are complete on Tuesday 27 October 2026. Explain why the target date is Friday 30 October. Henry explains the order check before pretend release.')
h('Round three in 15 minutes')
p('Pat says: “The PDF opens, but the CSV will not open.” Record this as an access problem. Do not mark receipt. Explain the approved recovery steps and next update. Then Pat confirms all three mock files open. Record receipt and ask about scope acceptance separately.')
h('The observer records what they saw')
p('Check: correct buyer; missing questions asked; 22-provider stop; clear date; matched simulated payment; separate reviewer; explicit release authority; all three files confirmed; owner and next date recorded.')
p('Observer name and date: __________________________________________')
p('Repeat this step: __________________________________________________')

page('A simple order handover sheet','Use this in practice. Real customer information belongs in the approved private register, not in this workbook or a public repository.')
question('Practice reference and enquiry received date and time')
question('What does the buyer sell and what decision do they want to make?')
question('Accepted scope version and buyer acceptance evidence')
p('Check the scope includes: area, provider types, factual rule/window, exclusions, grouping, quantity, source, files, recipient, fee/tax, terms and delivery date.')
question('Later of scope confirmation and source check; agreed calendar delivery date')
question('Source version and fingerprint; number of different providers; readiness evidence location')
question('Payment evidence and order match; unresolved payment issues')
question('Separate reviewer and result; exact final file fingerprints')
question('Approved recipient and delivery route; explicit release authority')
question('Access confirmation for all three files; acceptance; usefulness check date')
question('What is missing; next action; owner; next update date; time used')

page('Check what you have learned','Answer in your own words. Use short answers. The trainer should check the meaning, not the spelling.')
question('1. Who buys our Brief, and who appears in the shortlist?')
question('2. What are the three files the buyer receives?')
question('3. One provider has three locations. How many providers is that?')
question('4. We agreed 25 providers but only 22 qualify. What do you do?')
question('5. Scope is confirmed Tuesday 27 October; source check finishes Wednesday 28 October 2026. What is the target date?')
question('6. Payment arrives late. Does payment restart the clock? What happens next?')
question('7. Henry helped make the pack. Can he independently approve it?')
question('8. Pat says only the PDF opens. Can we record receipt? Why?')
question('9. A prospect asks us to stop marketing. What do we do?')
question('10. Does passing this quiz let us take payment or send a pack?')

page('Answers and practice signoff','Trainer use. Every question tests an important part of the job. Explain mistakes and repeat the affected step.')
table(['Question','A correct answer includes'],[
('1','A supplier buys the Brief. Care provider organisations are potential accounts in its shortlist.'),
('2','PDF report, CSV list, and source/licence manifest.'),
('3','One provider. Three locations are not three providers.'),
('4','Stop before payment. Get approval for a supported new scope or decline. Never pad.'),
('5','Monday 2 November 2026. Start from Wednesday; count Thursday, Friday and Monday.'),
('6','No. Verify payment before production. Keep the agreed date if it is still achievable. Otherwise Henry agrees a new written calendar date with the buyer before production.'),
('7','No. A separate competent person who did not make it must review it.'),
('8','No. The agreed recipient must confirm all three files open. Log and resolve the access problem.'),
('9','Record the request immediately in the do-not-contact system and stop marketing to that recipient.'),
('10','No. Training is separate from real readiness gates and order release authority.')],[.8,6.2])
h('Proposed learning standard')
p('Both learners answer all ten correctly and demonstrate the page 11 steps. If a step is wrong, explain it and repeat it. A supervisor witnesses the repeat. Do not write “trained” until this has actually happened.')
p('Sibongile completed learning check on: ______________________________')
p('Henry completed learning check on: _________________________________')
p('Witness name and date: ___________________________________________')
p('Steps to repeat and date: __________________________________________')
p('This signoff records learning only. It does not approve a commercial gate, real order or pack. A training witness does not automatically become the separate pack reviewer.',True)

page('Daily cards you can keep nearby','Keep service work inside the planned five-hour day. Reduce new prospecting when a paid order or recovery needs time.')
pic('day.png','Proposed production day. The six blocks total five hours. Breaks are outside these working hours. Training replaces other work; it is not added on top.')
h('Sibongile at the start of the day')
p('Read the approved inbox. Find today’s promises and blocked orders. For each active enquiry, check: next action, owner and update date. Confirm approved permissions before sending anything.')
h('Sibongile before finishing')
p('Record enquiry status, payment/review/access evidence and time used. Put the next date beside every open item. Tell Henry about anything that may miss its date. Leave a handover the backup can follow.')
h('Henry each day')
p('Check sources, payment problems and release decisions. Confirm production and separate review slots. Resolve blocked approvals explicitly. Cover the inbox only in the agreed backup slot.')
h('Friday and time zones')
p('The proposed Friday 14:00 to 15:00 UK meeting replaces an existing work hour. It is not extra work. Review open orders, time used, customer problems and next week’s capacity.')
p('Agree actual coverage times. If Sibongile works in South Africa, UK 09:00 is 10:00 there before 25 October 2026 and 11:00 there afterwards. Confirm her location rather than assuming it.')
p('Useful phrase: “Here is what happened, what I need checked, who owns the next step and when I will update you.”')

page('Henry checks readiness evidence','This page is the bridge to the detailed operating plan. The plain-language guide does not replace the full gate requirements.')
table(['Gate','Henry must locate the actual evidence'],[
('G0 People','Staff access, real inbox and register, confirmed backup and separate reviewer slots.'),
('G1 Source','Approved source/software version and fingerprints; fresh and complete enough for the scope; gaps resolved. Find the full technical source, comparison and timing checks in the operating plan.'),
('G2 Scope','Proven selection rules, supported provider set and different-provider count. Output matches what is being promised.'),
('G3 Marketing','Explicit campaign and recipient approval, relevant identity, privacy and do-not-contact checks, approved wording, and required phone screening.'),
('G4 Payment','Approved terms, tax and dates; actual merchant matching; proof for failures, mismatches, duplicates, cancellation, refunds and repeat processing.'),
('G5 Rehearsal','Full minimum 25-provider practice pack; separate review; real authorised controlled inbox and private access test; wrong-recipient denial, recovery and timing.'),
('G6 Each release','Accepted scope, matched payment, exact reviewed file fingerprints, authorised recipient, tested route and release authority.')],[1.5,5.5])
p('G0 to G5 must pass before a paid commitment or payment request. G6 must pass before each real release. Approved customer service messages use their approved route; G3 controls marketing. Keep all other applicable project gates in place.')
h('Evidence is more than a tick')
p('Record who approved, when, which exact source/files/system were checked, what the evidence showed and its limits. A planned test, fictional payment or discussion is not a successful real test. No decision by silence.')
p('The project review loop still applies. Grok investigates and retests. Codex applies the named fix. DeepSeek challenges a saved evidence pack that stays unchanged during review. Hermes sets the next agreed step. A producing model cannot approve its own work. This workbook does not complete that loop.')

page('Where this training comes from','Prepared for Sibongile and Henry from the consolidated CareGist plan and its checked feedback. Examples are fictional; schedules and learning standards are proposals.')
h('Main reference')
p('CareGist customer journey and 30 day operating plan, 11 October 2026. Read pages 4 to 9 for roles and the full order checks, pages 10 to 14 for capacity and dates, and pages 18 to 20 for draft messages and evidence.')
p('Local reference: artifacts/final-operating-plan/2026-10-11/CareGist_30_Day_Operating_Plan.pdf')
p('Source text fingerprint SHA256: 920792636f3aa9b5ffe23c60d8d373cfbeea85bf2f15f58efa0eab4cde472a6e')
h('Project rules')
p('AGENTS.md and .warroom/README.md and PIPELINE.md. Money-affecting and customer-facing actions need Henry’s explicit approval. Historical status entries are dated evidence; this training does not verify today’s live system.')
h('Primary guidance checked on 11 October 2026')
p('CQC using CQC data: explains provider/location data, rating levels, reuse acknowledgement and update limitations. https://www.cqc.org.uk/about-us/transparency/using-cqc-data')
p('ICO business to business marketing: public contact details are not automatic marketing permission. Rules depend on the recipient and use of personal data. Henry must use the current applicable guidance for a campaign. https://ico.org.uk/for-organisations/direct-marketing-and-privacy-and-electronic-communications/business-to-business-marketing/')
h('What this document proves')
p('The training materials are prepared. Completion needs an actual learning session and witnessed practice. Commercial readiness needs the actual operating evidence and required independent approvals. No real training attendance, purchase or delivery is claimed.')

page('Trainer cards for the pretend order','Read these aloud in round two. All names, approvals and records below are fictional. They are not proof that CareGist is ready to sell.')
h('Scope card')
p('Practice reference: TRAIN-001. Buyer: Pat at Practice Training Company. Objective: decide which organisations to investigate for staff training. Area: Birmingham local authority, used only as an example. Provider type: residential care for older adults, a training description that must match the approved source type in a real job.')
p('Rules: use active provider records at the pretend source date of 27 October 2026, with at least one matching location in the example area. Leave out already-known Practice Provider X. Count each provider ID once. Do not combine providers by commercial group. Quantity: 25 fictional providers after exclusions.')
p('Source: invented Practice Dataset V1. Files: PDF, CSV and source record, all version 1. Recipient: Pat, played by the trainer; no real email address. Fee: fictional £745 total, tax wording assumed approved for this exercise only. Terms: pretend practice terms V1, not real sale terms. Target release: Friday 30 October 2026.')
p('Pretend written reply from Pat on Tuesday 27 October: “I accept scope TRAIN-001 V1, its practice terms and the delivery date.” Source check: completed Tuesday 27 October. Trainer says readiness G0 to G5 is assumed for this paper exercise only.')
h('Payment card')
p('SIMULATION ONLY. Payment record TRAIN-PAY-001. Order TRAIN-001. Payer Practice Training Company. Status paid. Amount £745 GBP. Refund none. Dated Tuesday 27 October 2026. There is no real payment or merchant record.')
p('Fault to try next: change the payment amount to £700. Learner should stop production, tell Henry and resolve the mismatch before continuing. Restore the original card only after explaining that this is a new pretend scenario.')
h('Review and release card')
p('Sibongile and Henry are the pretend producers. Fictional Reviewer R did not make the pack. Pretend result: PASS for the three version 1 file cards on page 19. Fingerprint labels F-PDF-1, F-CSV-1 and F-SOURCE-1 are not real hashes. Real orders need real fingerprints and actual review evidence.')
p('Pretend Henry release authority on Friday 30 October: “TRAIN-001 V1 may be released to Pat through the assumed tested private route.” Nothing is sent. Ask the learner what changes if someone edits the CSV after review: stop and get the changed files checked again.')

page('Trainer file cards and buyer replies','Use these as pretend file descriptions. The exercise is spoken or on paper. These cards cannot replace opening actual files in the full rehearsal.')
h('PDF card version 1')
p('TRAINING ONLY. Label F-PDF-1. Research pack for TRAIN-001. It describes 25 invented providers, the agreed example area and rules. Summary: the list is for investigation; it does not prove a need for training or a future purchase. Sources and doubts are labelled. No real CQC facts are used.')
h('CSV card version 1')
p('TRAINING ONLY. Label F-CSV-1. Imagine 25 rows with provider labels P01 through P25, one different label per row. Each row has an invented matching location, factual fit reason, next qualification question and source label Practice Dataset V1. P01 through P25 are not real CQC IDs.')
p('Counting fault: imagine replacing P25 with another P01. The file still has 25 rows but only 24 different providers. It fails the agreed 25-provider scope. Stop and fix the selection before review or release.')
h('Source record card version 1')
p('TRAINING ONLY. Label F-SOURCE-1. Source: invented Practice Dataset V1. Pretend source date and collection date: 27 October 2026. Reuse: fictional exercise data only. It contains no real CQC data and grants no rights over any real source. File labels match the review card.')
h('Read these buyer replies one at a time')
p('Reply 1: “The PDF opens, but the CSV will not open.” Expected action: record the file and error, keep receipt unconfirmed, tell the owner and set a next update. Explain the authorised recovery route without sending anything.')
p('Reply 2, after pretend recovery on Friday 30 October: “I can now open the PDF, CSV and source record.” Expected action: record confirmed access on 30 October. Do not infer scope acceptance from this sentence.')
p('Reply 3: “The pack matches the scope we agreed.” Expected action: record product acceptance separately. Schedule the proposed usefulness check for Friday 6 November 2026.')
p('Observer asks: “What decision did this pack help Pat make?” Learner should ask Pat later. They must not invent a sale, saving or useful outcome.')

doc.save(OUT/'CareGist_Training_Sibongile_and_Henry.docx')
(OUT/'TRAINING.md').write_text('\n'.join(source))
print(json.dumps({'planned_pages':page_count,'words':len(' '.join(source).split()),'docx':str(OUT/'CareGist_Training_Sibongile_and_Henry.docx')}))
