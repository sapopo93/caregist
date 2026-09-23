"""Reproducible operator pack for the manual Territory Brief service.

No payment, messaging, or publication side effects. A draft is never sale approval.
Source observation timestamps are distinct from inspection/registration dates.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
from collections import Counter
from datetime import date, datetime, timedelta
from pathlib import Path

from api.services.pdf_writer import PdfBuilder

BUYERS = {'new_90', 'inadequate', 'requires_improvement', 'not_yet_inspected', 'stale_inspection'}
FIELDS = ['location_id', 'location_name', 'provider_id', 'provider_name', 'region',
          'local_authority', 'postcode', 'service_types', 'rating', 'registration_date',
          'inspection_date', 'observed_at', 'cqc_location_url', 'cqc_provider_url']
SHORTLIST_FIELDS = ['rank', 'provider_id', 'provider_name', 'matching_locations',
                    'representative_location_id', 'published_fact', 'reason',
                    'qualification_question', 'uncertainty', 'cqc_provider_url']


def _date(value):
    return date.fromisoformat(str(value)[:10]) if value else None


def qualifies(row: dict, buyer: str, as_of: date) -> bool:
    if buyer not in BUYERS:
        raise ValueError('Unknown buyer type')
    rating = str(row.get('overall_rating') or '').strip().lower()
    registered, inspected = _date(row.get('registration_date')), _date(row.get('last_inspection_date'))
    if buyer == 'new_90':
        return registered is not None and as_of - timedelta(days=90) <= registered <= as_of
    if buyer == 'inadequate':
        return rating == 'inadequate'
    if buyer == 'requires_improvement':
        return rating == 'requires improvement'
    if buyer == 'not_yet_inspected':
        return rating in ('', 'not yet inspected', 'no published rating')
    # Match the picker population, but describe missing dates truthfully below.
    threshold = as_of.replace(year=as_of.year - 3, day=min(as_of.day, 28)) if as_of.month == 2 else as_of.replace(year=as_of.year - 3)
    return inspected is None or inspected < threshold


def build_pack(rows: list[dict], *, region: str, buyer: str, service_type: str,
               observed_at: datetime, provider_names: dict[str, str],
               source_evidence: dict, order_reference: str, shortlist_target: int = 30) -> dict:
    if buyer not in BUYERS or not region.strip() or not order_reference.strip():
        raise ValueError('An exact region, supported buyer type and order reference are required')
    if not 25 <= shortlist_target <= 50:
        raise ValueError('Shortlist target must be between 25 and 50 organisations')
    if observed_at.tzinfo is None:
        raise ValueError('Observation timestamp must include a timezone')
    matched = []
    seen = set()
    for row in rows:
        if (str(row.get('status', '')).upper() != 'ACTIVE' or row.get('region') != region
                or not qualifies(row, buyer, observed_at.date())):
            continue
        services = str(row.get('service_types') or '')
        if service_type and service_type.casefold() not in services.casefold():
            continue
        lid, pid = str(row.get('id') or ''), str(row.get('provider_id') or '')
        if not lid or not pid or lid in seen:
            raise ValueError('Missing or duplicate source identifiers; repair source before export')
        seen.add(lid)
        matched.append({
            'location_id': lid, 'location_name': str(row.get('name') or ''),
            'provider_id': pid, 'provider_name': provider_names.get(pid, ''),
            'region': region, 'local_authority': row.get('local_authority') or '',
            'postcode': row.get('postcode') or '', 'service_types': services,
            'rating': row.get('overall_rating') or 'No published rating recorded',
            'registration_date': str(row.get('registration_date') or ''),
            'inspection_date': str(row.get('last_inspection_date') or ''),
            'observed_at': str(row.get('signal_checked_at') or row.get('last_updated') or ''),
            'cqc_location_url': f'https://www.cqc.org.uk/location/{lid}',
            'cqc_provider_url': f'https://www.cqc.org.uk/provider/{pid}',
        })
    matched.sort(key=lambda r: (r['provider_id'], r['location_id']))
    grouped = {}
    for row in matched:
        grouped.setdefault(row['provider_id'], []).append(row)
    # Research priority: larger matching location footprint, then stable CQC ID.
    # This is not an estimate of propensity to buy or regulatory risk.
    ordered = sorted(grouped.items(), key=lambda item: (-len(item[1]), item[0]))
    shortlist = []
    for pid, locations in ordered[:shortlist_target]:
        row = locations[0]
        fact = (f"Registration recorded on {row['registration_date']}" if buyer == 'new_90' else
                (f"Inspection date recorded as {row['inspection_date']}" if row['inspection_date'] else
                 'No inspection date is recorded in this source') if buyer == 'stale_inspection' else
                f"Recorded rating: {row['rating']}")
        shortlist.append({
            'rank': len(shortlist) + 1, 'provider_id': pid,
            'provider_name': provider_names.get(pid, ''), 'matching_locations': len(locations),
            'representative_location_id': row['location_id'], 'published_fact': fact,
            'reason': f'{len(locations)} active location(s) match the agreed region and buyer criterion. Review the linked locations before approaching the provider.',
            'qualification_question': 'Who owns supplier selection for these locations, and is the proposed service relevant to their current priorities?',
            'uncertainty': 'No evidence of buying intent, budget, vacancies or need for a specific supplier. Missing inspection/rating information does not prove a location has never been inspected.',
            'cqc_provider_url': row['cqc_provider_url'],
        })
    blockers = []
    reconciled = source_evidence.get('completed_at')
    if source_evidence.get('status') != 'completed' or not source_evidence.get('source_checksum_sha256') or not reconciled:
        blockers.append('Authoritative source reconciliation has not completed with a checksum')
    elif not 0 <= (observed_at - datetime.fromisoformat(str(reconciled))).total_seconds() <= 192 * 3600:
        blockers.append('Reconciliation is outside the 192-hour freshness window')
    if len(shortlist) < 25:
        blockers.append('Fewer than 25 distinct qualifying provider organisations')
    if any(not row['provider_name'] for row in shortlist):
        blockers.append('One or more shortlisted provider names lack source verification')
    if any(not row['observed_at'] for row in matched):
        blockers.append('One or more location observation timestamps are missing')
    return {'order_reference': order_reference, 'region': region, 'buyer_type': buyer,
            'service_type': service_type, 'extracted_at': observed_at.isoformat(),
            'source_evidence': source_evidence, 'blockers': blockers,
            'status': 'DRAFT - NOT FOR CUSTOMER DELIVERY',
            'locations': matched, 'shortlist': shortlist, 'provider_count': len(grouped)}


def _safe(value):
    text = str(value) if value is not None else ''
    return "'" + text if text.lstrip().startswith(('=', '+', '-', '@', '\t', '\r')) else text


def _csv(rows, fields):
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(fields)
    for row in rows:
        w.writerow([_safe(row.get(key, '')) for key in fields])
    return out.getvalue()


def render_summary(pack: dict) -> bytes:
    pdf = PdfBuilder(title='Territory Opportunity Brief', footer='CareGist | Rehearsal | Not for customer delivery')
    pdf.title_block(subtitle=pack['region'], meta_lines=[pack['status'], 'One-off price: GBP 745. H-Kay Limited is not VAT registered.'])
    pdf.heading('1. Scope and evidence')
    pdf.key_values([('Reference', pack['order_reference']), ('Buyer criterion', pack['buyer_type']),
                    ('Service type', pack['service_type'] or 'All'), ('Extracted at', pack['extracted_at']),
                    ('Qualifying locations', str(len(pack['locations']))),
                    ('Distinct providers', str(pack['provider_count'])), ('Shortlisted providers', str(len(pack['shortlist'])))])
    pdf.paragraph('This pack groups active CQC locations by provider ID. A location name is not substituted for its provider organisation name. The workbook and CSV preserve the location evidence behind each account.')
    pdf.paragraph('Extraction time describes this database read. Inspection and registration dates describe events. None of these dates alone is a complete source-reconciliation date.')
    pdf.heading('Release checks still required')
    pdf.bullets(pack['blockers'] + ['Operator checks every shortlist entry against its source.', 'Written scope, delivery date and applicable terms are agreed before payment.', 'Payment and receipt of all promised files are verified separately.'])
    pdf.page_break()
    pdf.heading('2. Territory structure and research priority')
    counts = Counter(r['local_authority'] or 'Not recorded' for r in pack['locations'])
    pdf.table(['Local authority', 'Matching locations'], [[k, str(v)] for k, v in counts.most_common(10)], widths=[4, 1.5])
    pdf.paragraph('Counts describe the selected buyer population, not the whole regional care market. Local authorities outside the ten largest remain in the dataset.')
    pdf.paragraph('Providers are ranked by their number of matching active locations, then CQC provider ID. Larger footprints warrant account-level research; this ordering does not predict purchases or claim commercial intent.')
    pdf.heading('Movement evidence')
    pdf.paragraph('This extraction is a point-in-time population. It does not establish rating changes or closures without a comparison to an earlier verified record. A current rating is not evidence of a recent downgrade.')
    pdf.page_break()
    pdf.heading('3. First accounts to research')
    for row in pack['shortlist'][:5]:
        pdf.heading(f"{row['rank']}. {row['provider_name'] or row['provider_id']}", level=2)
        pdf.paragraph(row['published_fact'] + '. ' + row['reason'])
        pdf.paragraph(row['cqc_provider_url'], muted=True)
    pdf.paragraph('The full shortlist is in the workbook and shortlist CSV. Each entry has a source link, stated fact, reason, qualification question and uncertainty.')
    pdf.page_break()
    pdf.heading('4. Recommended approach and limits')
    pdf.bullets(['Confirm each provider and representative location on CQC before contact.', 'Map the selected locations to the provider organisation in your CRM.', 'Check who selects suppliers and whether the proposed service fits their current priorities.', 'Record a reason to include or reject each account. Do not treat a rating as a claim of buying intent.'], ordered=True)
    pdf.heading('Files and definitions')
    pdf.paragraph('territory.csv contains qualifying care locations. shortlist.csv contains distinct provider accounts. territory.xlsx contains both tables and source notes. This four-page PDF summarises the scope, market structure, priorities and limits.')
    pdf.paragraph('A missing rating or inspection date means the source has no usable value. It does not establish that CQC has never inspected the location. Status, ownership and contacts can change after observation.')
    pdf.paragraph('CareGist is independent of the Care Quality Commission. Contains public sector information licensed under the Open Government Licence v3.0. https://www.nationalarchives.gov.uk/doc/open-government-licence/version/3/')
    pdf.paragraph('This rehearsal does not authorise sending this pack, taking payment or enabling automated checkout. A source-clean result is only one part of the release decision.')
    return pdf.build()


def write_pack(pack: dict, output: Path) -> dict:
    """Operator export. Uses the project's existing workbook dependency."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    output.mkdir(parents=True, exist_ok=False)
    (output / 'territory.csv').write_text(_csv(pack['locations'], FIELDS), encoding='utf-8-sig')
    (output / 'shortlist.csv').write_text(_csv(pack['shortlist'], SHORTLIST_FIELDS), encoding='utf-8-sig')
    (output / 'brief.pdf').write_bytes(render_summary(pack))
    wb = Workbook()
    notes = wb.active
    notes.title = 'Read first'
    for key, value in [('Status', pack['status']), ('Region', pack['region']), ('Buyer criterion', pack['buyer_type']),
                       ('Extracted at', pack['extracted_at']), ('Source evidence', json.dumps(pack['source_evidence'], default=str)),
                       ('Blockers', '; '.join(pack['blockers'])), ('Ranking', 'Matching location count descending, then provider ID'),
                       ('Licence', 'Contains public sector information licensed under the Open Government Licence v3.0')]:
        notes.append([key, _safe(value)])
    notes.column_dimensions['A'].width = 24
    notes.column_dimensions['B'].width = 100
    for title, rows, fields in [('Shortlist', pack['shortlist'], SHORTLIST_FIELDS), ('Territory locations', pack['locations'], FIELDS)]:
        ws = wb.create_sheet(title)
        ws.append(fields)
        for row in rows:
            ws.append([row.get(k, '') if isinstance(row.get(k), int) else _safe(row.get(k, '')) for k in fields])
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 30
        ws.row_dimensions[1].height = 32
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='245C4C')
            cell.alignment = Alignment(wrap_text=True)
    wb.save(output / 'territory.xlsx')
    evidence = {k: v for k, v in pack.items() if k not in ('locations', 'shortlist')}
    evidence['files'] = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in sorted(output.iterdir())}
    (output / 'evidence.json').write_text(json.dumps(evidence, indent=2, default=str) + '\n')
    return evidence
