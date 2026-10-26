"""Build the maintained Markdown manuals into an editable Word document.

Run with the Codex bundled Python runtime, then render with the documents
skill render_docx.py --emit_pdf. The renderer's PDF is the printable edition.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(__file__).resolve().parent
FILES = ['USER_MANUAL.md', 'TRAINING_MANUAL.md', 'SALES_GTM_MANUAL.md',
         'CRM_MANUAL.md', 'OPERATIONS_MANUAL.md']
REPO = 'https://github.com/sapopo93/caregist/blob/main/'


def inline(paragraph, text):
    pattern = r'\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*|`([^`]+)`'
    position = 0
    for match in re.finditer(pattern, text):
        paragraph.add_run(text[position:match.start()])
        if match.group(1):
            url = match.group(2)
            if not url.startswith(('https:', 'mailto:')):
                path = (BASE / url).resolve().relative_to(BASE.parent.parent)
                url = REPO + path.as_posix()
            relation = paragraph.part.relate_to(
                url,
                'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                is_external=True,
            )
            link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), relation)
            run = OxmlElement('w:r'); properties = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '174D6A'); properties.append(color)
            underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); properties.append(underline)
            run.append(properties); content = OxmlElement('w:t'); content.text = match.group(1)
            run.append(content); link.append(run); paragraph._p.append(link)
        else:
            run = paragraph.add_run(match.group(3) or match.group(4))
            run.bold = bool(match.group(3))
        position = match.end()
    paragraph.add_run(text[position:])


def add_table(doc, rows):
    values = [[cell.strip() for cell in row.strip().strip('|').split('|')] for row in rows]
    values = [row for row in values if not all(re.fullmatch(r'[:\- ]+', cell) for cell in row)]
    table = doc.add_table(rows=0, cols=len(values[0])); table.autofit = False
    widths = [1.15, 2.10, 3.15] if len(values[0]) == 3 else [2.2, 4.2]
    for column, width in zip(table.columns, widths): column.width = Inches(width)
    for index, row in enumerate(values):
        cells = table.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Inches(width)
            properties = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                border = OxmlElement('w:' + edge)
                for key, setting in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    border.set(qn('w:' + key), setting)
                borders.append(border)
            properties.append(borders)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EDF1' if index == 0 else ('F6F8FA' if index % 2 == 0 else 'FFFFFF')); properties.append(shade)
            margin = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                node = OxmlElement('w:' + edge); node.set(qn('w:w'), '90'); node.set(qn('w:type'), 'dxa'); margin.append(node)
            properties.append(margin)
            paragraph = cell.paragraphs[0]; paragraph.paragraph_format.space_after = Pt(3)
            inline(paragraph, value)
            for run in paragraph.runs: run.font.size = Pt(10); run.bold = index == 0 or run.bold
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); table.rows[-1]._tr.get_or_add_trPr().append(repeat)
        no_split = OxmlElement('w:cantSplit'); table.rows[-1]._tr.get_or_add_trPr().append(no_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def main():
    document = Document(); section = document.sections[0]
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    section.left_margin = section.right_margin = Inches(.72)
    section.top_margin = section.bottom_margin = Inches(.7)
    for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = document.styles[name]; style.font.name = 'Calibri'; style.font.color.rgb = RGBColor(0, 0, 0)
    for style in document.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = document.styles['Normal']; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7); normal.paragraph_format.line_spacing = 1.06
    document.styles['Title'].font.size = Pt(26)
    document.styles['Heading 1'].font.size = Pt(19)
    document.styles['Heading 2'].font.size = Pt(14)
    for name in ('Heading 1', 'Heading 2', 'Heading 3'):
        document.styles[name].paragraph_format.keep_with_next = True
    section.header.paragraphs[0].text = 'CareGist  |  Manuals  |  15 October 2026'
    footer = section.footer.paragraphs[0]; footer.text = 'Prepared instructions are separate from operating approvals  |  '
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for paragraph in (section.header.paragraphs[0], footer):
        for run in paragraph.runs: run.font.size = Pt(9)
    for file_index, filename in enumerate(FILES):
        if file_index: document.add_page_break()
        lines = (BASE / filename).read_text().splitlines(); index = 0
        while index < len(lines):
            line = lines[index]
            if not line.strip(): index += 1; continue
            if line.startswith('|'):
                rows = []
                while index < len(lines) and lines[index].startswith('|'):
                    rows.append(lines[index]); index += 1
                add_table(document, rows); continue
            if line.startswith('# '): inline(document.add_paragraph(style='Title'), line[2:])
            elif line.startswith('## '): inline(document.add_paragraph(style='Heading 1'), line[3:])
            elif re.match(r'^\d+\. ', line):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(.2)
                paragraph.paragraph_format.first_line_indent = Inches(-.2)
                inline(paragraph, line)
            elif line.startswith('- '): inline(document.add_paragraph(style='List Bullet'), line[2:])
            else: inline(document.add_paragraph(), line)
            index += 1
    document.core_properties.title = 'CareGist user training sales CRM and operations manuals'
    document.core_properties.author = 'CareGist'
    output = BASE / 'CareGist_Manuals.docx'; document.save(output)
    print(output)


if __name__ == '__main__': main()
